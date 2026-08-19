"""Contracts for slide partname allocation.

Deriving the next slide's name from the slide count collides whenever the partname
sequence has a gap. This is an upstream defect, not a paper-introduced one: stock
python-pptx 1.0.2 reaches it by opening any deck whose numbering has a gap, writes
duplicate zip members, and reopens the result without error while silently dropping a
slide. `Slides.delete` (a paper addition) only made it easier to reach in-process.
"""

from __future__ import annotations

import pytest

from pptx import Presentation
from pptx.errors import PackageLimitError
from pptx.util import Inches

from .contract import save_reopen


def _deck(slide_count=5):
    prs = Presentation()
    layout = prs.slide_layouts[6]
    for idx in range(slide_count):
        slide = prs.slides.add_slide(layout)
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        shape.text_frame.text = "SLIDE-%d" % (idx + 1)
    return prs


def _slide_partnames(prs):
    return sorted(
        str(part.partname)
        for part in prs.part.package.iter_parts()
        if "/slides/slide" in str(part.partname)
    )


def _texts(prs):
    return [
        next(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame) for slide in prs.slides
    ]


def test_add_after_delete_reuses_the_freed_partname_not_a_live_one():
    """The gap left by the delete is what the next add must take."""
    prs = _deck()
    prs.slides.delete(prs.slides[1])  # frees /ppt/slides/slide2.xml
    assert "/ppt/slides/slide2.xml" not in _slide_partnames(prs)

    prs.slides.add_slide(prs.slide_layouts[6])

    partnames = _slide_partnames(prs)
    assert "/ppt/slides/slide2.xml" in partnames
    assert len(partnames) == len(set(partnames))


def test_deck_survives_save_reopen_after_delete_then_add():
    """The whole point: the written package must be readable."""
    prs = _deck()
    prs.slides.delete(prs.slides[1])
    new_slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = new_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    shape.text_frame.text = "NEW"

    assert _texts(save_reopen(prs)) == ["SLIDE-1", "SLIDE-3", "SLIDE-4", "SLIDE-5", "NEW"]


def test_repeated_delete_and_add_never_collides():
    """Churn is where a count-based allocator drifts furthest from the free set."""
    prs = _deck(slide_count=6)

    for _ in range(4):
        prs.slides.delete(prs.slides[0])
        prs.slides.add_slide(prs.slide_layouts[6])
        partnames = _slide_partnames(prs)
        assert len(partnames) == len(set(partnames)), partnames

    assert len(save_reopen(prs).slides) == 6


def test_add_slide_is_safe_on_a_gapped_sequence_not_produced_by_delete():
    """The upstream-reachable case: a gap that came from the file, not from `delete`.

    Numbering is not required to be contiguous, and tools that remove a slide without
    renumbering produce exactly this. Stock python-pptx collides here and writes a
    package that silently loses a slide.
    """
    from pptx.opc.packuri import PackURI

    prs = _deck(slide_count=4)
    for part in list(prs.part.package.iter_parts()):
        if str(part.partname) == "/ppt/slides/slide4.xml":
            part.partname = PackURI("/ppt/slides/slide5.xml")
    assert _slide_partnames(prs) == [
        "/ppt/slides/slide1.xml",
        "/ppt/slides/slide2.xml",
        "/ppt/slides/slide3.xml",
        "/ppt/slides/slide5.xml",
    ]

    new_slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = new_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    shape.text_frame.text = "NEW"

    partnames = _slide_partnames(prs)
    assert len(partnames) == len(set(partnames))
    assert _texts(save_reopen(prs)) == ["SLIDE-1", "SLIDE-2", "SLIDE-3", "SLIDE-4", "NEW"]


def test_writer_refuses_to_serialize_two_parts_sharing_a_partname(tmp_path):
    """Defence in depth: duplicate detection existed only on the reading side.

    A package whose parts collide serializes into a zip with duplicate member names;
    readers keep the last copy and the other part vanishes silently.
    """
    prs = _deck(slide_count=3)
    slide_parts = [
        part for part in prs.part.package.iter_parts() if "/slides/slide" in str(part.partname)
    ]
    slide_parts[1].partname = slide_parts[0].partname
    destination = tmp_path / "collision.pptx"

    with pytest.raises(PackageLimitError, match="sharing a partname"):
        prs.save(str(destination))

    assert not destination.exists()
