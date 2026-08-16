"""The round-trip covenant at scale: anything paper-pptx can write, it can read.

`pptx._zipguard` once enforced six numeric resource ceilings on package intake, so a deck this
library saved without complaint could be refused when reopened: 2,100 slides is 4,236 ZIP
members, and the 4,096-member ceiling refused exactly that file. Those ceilings are gone, and
these tests are the standing proof -- one per ceiling, at the scale that used to refuse it.

The invariant is general, not a threshold: whatever `save()` writes, `Presentation()` reads.
Every case asserts its own PRECONDITION, that the artifact really does exceed the deleted
ceiling, because a test that quietly stops being large is worse than no test.

The upstream-parity half of the change -- an oversized non-`.pptx` path raising
`PackageNotFoundError` rather than `PackageLimitError` -- is covered by
`test_package_io_hardening.py::test_huge_non_zip_path_raises_package_not_found` and is not
duplicated here. The surviving structural refusals are pinned in `test_zipguard_structural.py`.

Opt-in: the `big_io` cases write and read a package larger than 256 MiB and are skipped by
default. Enable them with `PAPER_BIG_IO=1 uv run pytest tests/paper` or
`uv run pytest tests/paper -m big_io`.

Runtime of this module: about 8 s by default, about 10 s with `big_io` enabled.
"""

from __future__ import annotations

import io
import random
import re
import zipfile

import pytest
from PIL import Image

from pptx import Presentation, _zipguard
from pptx.opc import serialized
from pptx.util import Inches

from . import contract

_MIB = 1 << 20

# -- the deleted ceilings, kept as the scale each case must exceed. Their only remaining
# -- purpose is to be surpassed; nothing in `src/` defines them anymore.
_OLD_MEMBER_COUNT_CEILING = 4_096
_OLD_XML_MEMBER_CEILING = 64 * _MIB
_OLD_BINARY_MEMBER_CEILING = 256 * _MIB
_OLD_TOTAL_EXPANDED_CEILING = 512 * _MIB
_OLD_COMPRESSED_CEILING = 256 * _MIB

_BIG_MEMBER_COUNT_DECK = 2_100
_IMAGE_COUNT = 6

# -- ~192 KiB of highly compressible filler: lets a member expand past a ceiling without
# -- costing the disk (deflated) or the wall clock (stored).
_FILLER = b"<row>expansion accounting placeholder row</row>" * 4096


def _deck(slide_count: int) -> Presentation:
    """Return a presentation of `slide_count` blank slides, each carrying one textbox.

    Two ZIP members per slide (the slide part and its rels item), so 2,100 slides clears the
    old 4,096-member ceiling.
    """
    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    for index in range(slide_count):
        slide = presentation.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text_frame.text = "slide %d" % index
    return presentation


def _noise_png(seed: int) -> bytes:
    """Return a distinct, incompressible 64x64 PNG.

    Distinctness is load-bearing: python-pptx hashes image bytes and collapses duplicates into
    a single media part, so repeating one image would add no bytes and measure nothing.
    """
    pixels = random.Random(seed).randbytes(64 * 64 * 3)
    buf = io.BytesIO()
    Image.frombytes("RGB", (64, 64), pixels).save(buf, format="PNG")
    return buf.getvalue()


def _build_zip(path, names, member_bytes: int, compress: int = zipfile.ZIP_DEFLATED) -> None:
    """Write a ZIP at `path` holding one `member_bytes`-long member per name in `names`.

    Built directly rather than through `Presentation.save()` because these sizes are only
    reachable as raw archive shapes; the read path under test is the same either way.
    """
    with zipfile.ZipFile(path, "w", compress, compresslevel=1) as archive:
        for name in names:
            with archive.open(name, "w") as member:
                written = 0
                while written < member_bytes:
                    block = _FILLER[: member_bytes - written]
                    member.write(block)
                    written += len(block)


def _guarded_open(path):
    """Open `path` through real guarded package intake; return its member names.

    This is `preflight_zip` plus `GuardedZipReader` over every member -- the exact read path
    the deleted ceilings gated -- without requiring a well-formed OPC part graph.
    """
    return serialized._PhysPkgReader.factory(str(path)).partnames


# ------------------------------------------------------------------- round-trip invariant


@pytest.mark.parametrize("slide_count", [1, 100, _BIG_MEMBER_COUNT_DECK])
def test_saved_deck_reopens_at_any_slide_count(slide_count):
    reopened = contract.save_reopen(_deck(slide_count))

    assert len(reopened.slides) == slide_count


def test_saved_deck_of_distinct_images_reopens():
    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    for seed in range(_IMAGE_COUNT):
        slide = presentation.slides.add_slide(blank)
        slide.shapes.add_picture(io.BytesIO(_noise_png(seed)), Inches(1), Inches(1))
    media = sorted(
        name
        for name in contract.zip_member_map(contract.save_to_bytes(presentation))
        if name.startswith("ppt/media/")
    )
    assert len(media) == _IMAGE_COUNT, "images collapsed; they are not distinct after all"

    reopened = contract.save_reopen(presentation)

    assert len(reopened.slides) == _IMAGE_COUNT
    assert media == sorted(
        name
        for name in contract.zip_member_map(contract.save_to_bytes(reopened))
        if name.startswith("ppt/media/")
    )


# ------------------------------------------------------------- one case per deleted ceiling


def test_saved_deck_past_the_member_count_ceiling_reopens():
    """The ceiling that actually shipped broken: this deck saved fine and refused to reopen.

    Also the central-directory case: a 4,236-member directory is far past trivial, though it
    stays well under the 16 MiB that ceiling named -- no cheap artifact reaches that, and the
    same arithmetic is deleted either way.
    """
    blob = contract.save_to_bytes(_deck(_BIG_MEMBER_COUNT_DECK))
    members = contract.zip_member_map(blob)
    assert len(members) > _OLD_MEMBER_COUNT_CEILING, (
        "fixture lost its bite; %d slides no longer exceed %d ZIP members"
        % (_BIG_MEMBER_COUNT_DECK, _OLD_MEMBER_COUNT_CEILING)
    )
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        central_directory_bytes = len(blob) - archive.start_dir
    assert central_directory_bytes > 128 * 1024

    reopened = Presentation(io.BytesIO(blob))

    assert len(reopened.slides) == _BIG_MEMBER_COUNT_DECK


def test_package_past_the_xml_member_ceiling_opens(tmp_path):
    """One XML part expanding past 64 MiB -- ~0.5 MiB on disk, because it is repetitive."""
    target = tmp_path / "big-part.pptx"
    _build_zip(target, ["ppt/slides/slide1.xml"], 68 * _MIB)
    with zipfile.ZipFile(target) as archive:
        (info,) = archive.infolist()
    assert info.file_size > _OLD_XML_MEMBER_CEILING

    assert len(_guarded_open(target)) == 1


def test_package_past_the_expanded_total_ceiling_opens(tmp_path):
    """540 MiB expanded across nine members, none individually past the per-member ceiling.

    Keeping every member under 64 MiB isolates the total from the per-member ceiling: this
    package was refused for its sum alone. Compressible XML keeps it a few MiB on disk.
    """
    target = tmp_path / "expanded.pptx"
    names = ["ppt/slides/slide%d.xml" % index for index in range(9)]
    _build_zip(target, names, 60 * _MIB)
    with zipfile.ZipFile(target) as archive:
        infos = archive.infolist()
    assert max(info.file_size for info in infos) < _OLD_XML_MEMBER_CEILING
    assert sum(info.file_size for info in infos) > _OLD_TOTAL_EXPANDED_CEILING

    assert len(_guarded_open(target)) == len(names)


@pytest.mark.big_io
def test_package_past_the_compressed_and_binary_member_ceilings_opens(tmp_path):
    """264 MiB of genuinely-on-disk bytes in one member: the two ceilings that cost real I/O.

    Stored rather than deflated, so the archive really is that large on disk (the compressed
    ceiling) and its single member really does expand past 256 MiB (the binary-member ceiling).
    """
    target = tmp_path / "huge.pptx"
    _build_zip(target, ["ppt/media/image1.bin"], 264 * _MIB, compress=zipfile.ZIP_STORED)
    assert target.stat().st_size > _OLD_COMPRESSED_CEILING
    with zipfile.ZipFile(target) as archive:
        (info,) = archive.infolist()
    assert info.file_size > _OLD_BINARY_MEMBER_CEILING

    assert len(_guarded_open(target)) == 1


# ------------------------------------------------------------------------ directory entries


def _rezip_with_directory_entries(source, destination) -> int:
    """Rewrite `source` the way an unzip-edit-rezip loop does; return the folder count.

    The folder records are emitted *interleaved* with the members, ahead of the first member
    in each directory, which is what `zip -r` and `shutil.make_archive` produce. Writing them
    all at the front instead would leave every member's data region contiguous, and a reader
    that mishandled them could still pass.
    """
    written = 0
    with zipfile.ZipFile(source) as incoming, zipfile.ZipFile(
        destination, "w", zipfile.ZIP_DEFLATED
    ) as outgoing:
        seen: set[str] = set()
        for info in incoming.infolist():
            folder = ""
            for segment in info.filename.split("/")[:-1]:
                folder += segment + "/"
                if folder in seen:
                    continue
                seen.add(folder)
                record = zipfile.ZipInfo(folder)
                record.external_attr = (0o40755 << 16) | 0x10
                outgoing.writestr(record, b"")
                written += 1
            outgoing.writestr(info.filename, incoming.read(info.filename))
    return written


def test_deck_rezipped_with_directory_entries_opens_and_round_trips(tmp_path):
    """A ZIP directory record is not an OPC part, so it is ignored rather than refused.

    PowerPoint opens such a package, so refusing it would reject decks the application
    accepts. Every unzip-edit-rezip pipeline emits these records, which is the common way
    a `.pptx` gets patched by hand.
    """
    original = tmp_path / "original.pptx"
    _deck(2).save(str(original))
    rezipped = tmp_path / "rezipped.pptx"
    folder_count = _rezip_with_directory_entries(original, rezipped)
    assert folder_count > 0, "fixture built no directory records"

    reopened = Presentation(str(rezipped))

    assert len(reopened.slides) == 2
    saved = tmp_path / "saved.pptx"
    reopened.save(str(saved))
    with zipfile.ZipFile(saved) as archive:
        names = archive.namelist()
    # -- dropped on the way out, exactly as upstream and PowerPoint drop them
    assert [name for name in names if name.endswith("/")] == []
    assert len(Presentation(str(saved)).slides) == 2


def test_directory_entries_do_not_trip_the_content_type_refusal(tmp_path):
    """A folder record resolves to no content type, and must not be asked to.

    It has no extension for a `Default` to match and no `Override` naming it. Checking
    physical records rather than part records would refuse the package here instead.
    """
    original = tmp_path / "original.pptx"
    _deck(1).save(str(original))
    rezipped = tmp_path / "rezipped.pptx"
    _rezip_with_directory_entries(original, rezipped)

    partnames = _guarded_open(rezipped)

    assert not [name for name in partnames if str(name).endswith("/")]
    assert len(Presentation(str(rezipped)).slides) == 1


# --------------------------------------------------------------------- guard against regrowth


def test_zipguard_exposes_no_policy_ceiling_constant():
    """`_MAX_END_COMMENT_BYTES` is the ZIP *format*'s 16-bit comment-length field maximum.

    It is a structural fact about the container, not a policy about how much a caller may
    open, so it is the one survivor. Asserted by introspecting the imported module rather than
    grepping its source: what matters is that no ceiling is reachable at runtime.
    """
    ceilings = sorted(name for name in vars(_zipguard) if re.fullmatch(r"_?MAX_\w+", name))

    assert ceilings == ["_MAX_END_COMMENT_BYTES"]
