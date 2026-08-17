"""Adversarial coverage for guarded package reads and atomic ordinary saves."""

from __future__ import annotations

import io
import os
import stat
import warnings
import zipfile

import pytest
from lxml import etree

from pptx import Presentation
from pptx.errors import PackageLimitError, PaperRefusal, UnsupportedStructureError
from pptx.exc import PackageNotFoundError
from pptx.opc import serialized

from . import corpus


def _minimal_path():
    return corpus.fixture_path("self_generated/minimal_clean.pptx")


def test_normal_open_refuses_duplicate_members(tmp_path):
    target = tmp_path / "duplicate.pptx"
    source = _minimal_path()
    with zipfile.ZipFile(source) as incoming, zipfile.ZipFile(target, "w") as outgoing:
        for info in incoming.infolist():
            outgoing.writestr(info, incoming.read(info.filename))
        slide = incoming.read("ppt/slides/slide1.xml")
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            outgoing.writestr("ppt/slides/slide1.xml", slide)

    assert issubclass(PackageLimitError, PaperRefusal)
    with pytest.raises(PackageLimitError, match="duplicate member"):
        Presentation(target)


def test_normal_open_refuses_noncanonical_member_names(tmp_path):
    target = tmp_path / "noncanonical.pptx"
    with zipfile.ZipFile(target, "w") as archive:
        archive.writestr("../ppt/presentation.xml", b"<presentation/>")

    with pytest.raises(PackageLimitError, match="noncanonical"):
        Presentation(target)


def test_saved_repetitive_deck_reopens(tmp_path):
    """Regression: the save -> reopen covenant must hold for this package's OWN output.

    Machine-generated decks (thousands of near-identical paragraphs) legitimately
    exceed any expanded-to-compressed ratio a hostile archive would need. A ratio guard once
    refused such files at reopen. No resource ceiling gates intake now: whatever this
    package writes, it reads. `test_package_round_trip.py` asserts that covenant at the
    scales the deleted numeric ceilings used to refuse.
    """
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(4)).text_frame
    for index in range(12_000):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = "Quarterly revenue synergy alignment placeholder row"
    target = tmp_path / "repetitive.pptx"
    presentation.save(target)

    with zipfile.ZipFile(target) as archive:
        worst_ratio = max(
            info.file_size / info.compress_size
            for info in archive.infolist()
            if info.compress_size
        )
    assert worst_ratio > 100, "fixture lost its bite; deck no longer compresses past 100:1"

    reopened = Presentation(target)
    assert len(reopened.slides) == 1


def _rewrite_package(source, target, transform):
    with zipfile.ZipFile(source) as incoming, zipfile.ZipFile(target, "w") as outgoing:
        for info in incoming.infolist():
            replacement = transform(info.filename, incoming.read(info.filename))
            if replacement is not None:
                outgoing.writestr(info, replacement)


def test_normal_open_refuses_a_missing_relationship_target(tmp_path):
    target = tmp_path / "missing-target.pptx"

    def transform(name, blob):
        if name == "_rels/.rels":
            return blob.replace(b"ppt/presentation.xml", b"ppt/missing-part.xml")
        return blob

    _rewrite_package(_minimal_path(), target, transform)

    with pytest.raises(UnsupportedStructureError, match="targets missing part"):
        Presentation(target)


def test_normal_open_accepts_an_unreachable_part_and_drops_it_on_save(tmp_path):
    """PowerPoint opens this package and drops the orphan on its own next save.

    Refusing to open it would refuse a deck PowerPoint accepts, so `save()` matching that
    behaviour is the whole of the contract. `.bin` is a declared Default here, so the part
    has a content type; one without a declared type is refused, in
    `test_content_type_coverage.py`.
    """
    target = tmp_path / "unreachable.pptx"
    _rewrite_package(_minimal_path(), target, lambda _name, blob: blob)
    with zipfile.ZipFile(target, "a") as archive:
        archive.writestr("ppt/orphan.bin", b"would be dropped on save")

    presentation = Presentation(target)
    assert len(presentation.slides) == len(Presentation(_minimal_path()).slides)

    saved = tmp_path / "resaved.pptx"
    presentation.save(saved)
    with zipfile.ZipFile(saved) as archive:
        assert "ppt/orphan.bin" not in archive.namelist()
    Presentation(saved)


def test_normal_open_refuses_duplicate_relationship_ids(tmp_path):
    target = tmp_path / "duplicate-rid.pptx"

    def transform(name, blob):
        if name != "_rels/.rels":
            return blob
        root = etree.fromstring(blob)
        root.append(etree.fromstring(etree.tostring(root[0])))
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    _rewrite_package(_minimal_path(), target, transform)

    with pytest.raises(UnsupportedStructureError, match="duplicate ids"):
        Presentation(target)


def _duplicate_first_rel(retarget: bool):
    """Return a transform duplicating the first relationship, optionally to a new target."""

    def transform(name, blob):
        if name != "_rels/.rels":
            return blob
        root = etree.fromstring(blob)
        clone = etree.fromstring(etree.tostring(root[0]))
        if retarget:
            clone.set("Target", "docProps/core.xml")
        root.append(clone)
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    return transform


def test_duplicate_relationship_id_message_distinguishes_redundant_from_conflicting(tmp_path):
    """Both are refused -- PowerPoint refuses either -- but the repair differs.

    Two declarations naming the same target are redundant: delete one and the meaning is
    unchanged. Two naming different targets leave the id with no single meaning, so the
    caller has to decide which was intended. The message says which it found.
    """
    redundant = tmp_path / "redundant.pptx"
    _rewrite_package(_minimal_path(), redundant, _duplicate_first_rel(retarget=False))
    with pytest.raises(UnsupportedStructureError, match="one declaration is redundant"):
        Presentation(redundant)

    conflicting = tmp_path / "conflicting.pptx"
    _rewrite_package(_minimal_path(), conflicting, _duplicate_first_rel(retarget=True))
    with pytest.raises(UnsupportedStructureError, match="no single meaning"):
        Presentation(conflicting)


def test_relationship_refusals_name_a_remedy(tmp_path):
    """A refusal exists so a caller can repair the package without reading library code.

    Both refuse input PowerPoint also refuses, so the caller's next step is the whole
    value of the message.
    """
    duplicate = tmp_path / "dup.pptx"
    _rewrite_package(_minimal_path(), duplicate, _duplicate_first_rel(retarget=False))
    with pytest.raises(UnsupportedStructureError, match="Remove the extra declaration"):
        Presentation(duplicate)

    missing = tmp_path / "missing.pptx"

    def drop_target(name, blob):
        if name == "_rels/.rels":
            return blob.replace(b"ppt/presentation.xml", b"ppt/missing-part.xml")
        return blob

    _rewrite_package(_minimal_path(), missing, drop_target)
    with pytest.raises(UnsupportedStructureError, match="Recover the missing part"):
        Presentation(missing)


def test_path_save_failure_preserves_existing_file_and_mode(tmp_path, monkeypatch):
    presentation = Presentation(_minimal_path())
    destination = tmp_path / "existing.pptx"
    destination.write_bytes(b"known-good destination")
    destination.chmod(0o640)
    original_write = serialized._ZipPkgWriter.write
    writes = 0

    def fail_during_write(self, pack_uri, blob):
        nonlocal writes
        writes += 1
        if writes == 2:
            raise OSError("forced ZIP write failure")
        return original_write(self, pack_uri, blob)

    monkeypatch.setattr(serialized._ZipPkgWriter, "write", fail_during_write)
    with pytest.raises(OSError, match="forced ZIP write failure"):
        presentation.save(destination)

    assert destination.read_bytes() == b"known-good destination"
    assert stat.S_IMODE(destination.stat().st_mode) == 0o640
    assert not list(tmp_path.glob(".existing.pptx.*.partial"))


def test_stream_save_failure_restores_existing_bytes_and_position():
    class FailingOnceStream(io.BytesIO):
        def __init__(self, initial):
            super().__init__(initial)
            self._fail_next_write = True

        def write(self, data):
            if self._fail_next_write:
                self._fail_next_write = False
                super().write(data[: min(16, len(data))])
                raise OSError("forced destination write failure")
            return super().write(data)

    presentation = Presentation(_minimal_path())
    original = b"ORIGINAL STREAM CONTENT"
    destination = FailingOnceStream(original)
    destination.seek(7)

    with pytest.raises(OSError, match="forced destination write failure"):
        presentation.save(destination)

    assert destination.getvalue() == original
    assert destination.tell() == 7


def test_huge_non_zip_path_raises_package_not_found(tmp_path, monkeypatch):
    """Upstream parity: size never gates intake; a non-zip path is simply not a package."""
    target = tmp_path / "huge.bin"
    target.write_bytes(b"not a zip")
    monkeypatch.setattr(os.path, "getsize", lambda _path: 1 << 40)

    with pytest.raises(PackageNotFoundError):
        Presentation(target)


def test_unsnapshottable_stream_is_written_rather_than_refused():
    """A destination that cannot be read back still receives the package.

    Rollback needs to read the destination, so it is offered where reading is possible
    rather than required before saving at all. Requiring it refused `open(path, "wb")`,
    sockets and pipes, which upstream accepted. Those callers get upstream's guarantee:
    a fully-serialized package is written, and nothing is taken back if the copy fails.
    """

    class UnreadableDestination(io.BytesIO):
        def read(self, size=-1):
            raise OSError("forced snapshot read failure")

    presentation = Presentation(_minimal_path())
    destination = UnreadableDestination(b"ORIGINAL STREAM CONTENT")
    destination.seek(7)

    presentation.save(destination)

    written = destination.getvalue()
    assert written.startswith(b"ORIGINA"), "the package was not written at the stream position"
    assert zipfile.ZipFile(io.BytesIO(written[7:])).testzip() is None


def test_successful_path_save_preserves_existing_mode(tmp_path):
    presentation = Presentation(_minimal_path())
    destination = tmp_path / "existing.pptx"
    destination.write_bytes(b"old")
    destination.chmod(0o604)

    presentation.save(destination)

    assert stat.S_IMODE(destination.stat().st_mode) == 0o604
    assert len(Presentation(io.BytesIO(destination.read_bytes())).slides) == 1


def test_new_path_save_honors_umask(tmp_path):
    """Regression: a save to a NEW path must not keep mkstemp's private 0600 mode."""
    presentation = Presentation(_minimal_path())
    destination = tmp_path / "brand-new.pptx"
    previous_umask = os.umask(0o027)
    try:
        presentation.save(destination)
    finally:
        os.umask(previous_umask)

    assert stat.S_IMODE(destination.stat().st_mode) == 0o666 & ~0o027
    assert len(Presentation(io.BytesIO(destination.read_bytes())).slides) == 1


def test_patch_save_honors_umask_and_existing_mode(tmp_path):
    """`patch_save` output modes follow the same contract as ordinary path saves."""
    from pptx.package import patch_save

    source = _minimal_path()
    destination = tmp_path / "patched.pptx"
    previous_umask = os.umask(0o027)
    try:
        patch_save(str(source), Presentation(source), str(destination))
    finally:
        os.umask(previous_umask)
    assert stat.S_IMODE(destination.stat().st_mode) == 0o666 & ~0o027

    destination.chmod(0o604)
    patch_save(str(source), Presentation(source), str(destination))
    assert stat.S_IMODE(destination.stat().st_mode) == 0o604
