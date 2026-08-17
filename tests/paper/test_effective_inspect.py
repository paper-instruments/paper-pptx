"""Contract tests: effective-style inspection (`pptx.inspect`).

Sidecar-driven against the branded-template and clrMap fixtures, determinism-goldened, and
independently cross-checked against LibreOffice's own resolution of the same source deck.
The API is read-only: proven here by part-snapshot equality around every call.
"""

from __future__ import annotations

import hashlib
import json

import pytest
from lxml import etree

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.errors import UnsupportedStructureError
from pptx.inspect import (
    BULLET_FOLLOWS_TEXT,
    content_hash,
    effective_font,
    effective_paragraph_format,
    inspect_text,
)
from pptx.util import Emu

from . import corpus
from .contract import save_reopen, snapshot_parts

BRANDED = "self_generated/branded_template.pptx"
CLRMAP = "self_generated/clrmap_remap.pptx"
GAUNTLET = "self_generated/gauntlet.pptx"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _open(relpath):
    return Presentation(str(corpus.fixture_path(relpath)))


def _ground_truth(relpath):
    return corpus.load_sidecar(relpath)["ground_truth"]


def _supplied_levels(effective_value):
    return [step.level for step in effective_value.provenance if step.supplied]


# ------------------------------------------------- sidecar-driven resolution (branded deck)


def test_title_size_resolves_through_layout_override():
    expected = _ground_truth(BRANDED)["expected_effective"]["title_run"]
    font = _open(BRANDED).slides[0].shapes.title.text_frame.paragraphs[0].runs[0].effective_font()
    assert font.size.resolved
    assert font.size.value_pt == expected["size_centipoints"] / 100.0
    assert font.size.value == Emu(int(expected["size_centipoints"] * 127))
    assert _supplied_levels(font.size) == ["layout placeholder lstStyle lvl1"]


def test_placeholder_inheritance_requires_an_exact_layout_idx_match():
    prs = _open(BRANDED)
    title = prs.slides[0].shapes.title
    title.element.ph.idx = 999

    with pytest.raises(UnsupportedStructureError, match="requires an exact idx match"):
        title.text_frame.paragraphs[0].runs[0].effective_font()


def test_layout_color_map_override_applies_to_layout_inherited_color():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    paragraph = title.text_frame.paragraphs[0]
    paragraph.clear()
    run = paragraph.add_run()
    run.text = "Layout-mapped color"
    layout = slide.slide_layout
    layout_title = layout.placeholders.get(idx=title.placeholder_format.idx)
    lst_style = layout_title._element.find("{%s}txBody/{%s}lstStyle" % (_P, _A))
    level_properties = etree.SubElement(lst_style, "{%s}lvl1pPr" % _A)
    default_run_properties = etree.SubElement(level_properties, "{%s}defRPr" % _A)
    solid_fill = etree.SubElement(default_run_properties, "{%s}solidFill" % _A)
    etree.SubElement(solid_fill, "{%s}schemeClr" % _A).set("val", "tx1")
    clr_map_override = layout._element.find("{%s}clrMapOvr" % _P)
    if clr_map_override is None:
        clr_map_override = etree.SubElement(layout._element, "{%s}clrMapOvr" % _P)
    else:
        clr_map_override.clear()
    override = etree.SubElement(clr_map_override, "{%s}overrideClrMapping" % _A)
    mapping = dict(prs.slide_masters[0].element.find("{%s}clrMap" % _P).attrib)
    mapping["tx1"] = "accent3"
    for key, value in mapping.items():
        override.set(key, value)

    effective = run.effective_font().color_rgb

    assert effective.value == "9BBB59"
    assert any(step.level == "layout clrMapOvr" for step in effective.provenance)


def test_title_name_resolves_through_master_theme_reference():
    expected = _ground_truth(BRANDED)["expected_effective"]["title_run"]
    font = _open(BRANDED).slides[0].shapes.title.text_frame.paragraphs[0].runs[0].effective_font()
    assert font.name.resolved
    assert font.name.value == expected["font_name"]
    assert _supplied_levels(font.name) == ["theme fontScheme majorFont"]
    # -- the +mj-lt reference itself is recorded on the master step of the chain
    master_steps = [s for s in font.name.provenance if "txStyles titleStyle" in s.level]
    assert master_steps
    assert "+mj-lt" in master_steps[0].detail


def test_body_level_one_resolves_from_master_body_style():
    expected = _ground_truth(BRANDED)["expected_effective"]["body_paragraph_0_run"]
    font = (
        _open(BRANDED).slides[0].placeholders[1].text_frame.paragraphs[0].runs[0].effective_font()
    )
    assert font.size.value_pt == expected["size_centipoints"] / 100.0
    assert font.name.value == expected["font_name"]  # -- "Trebuchet MS", explicit at master
    assert _supplied_levels(font.size) == ["master txStyles bodyStyle lvl1"]
    assert _supplied_levels(font.name) == ["master txStyles bodyStyle lvl1"]


def test_body_level_two_resolves_size_from_master_and_name_from_theme():
    expected = _ground_truth(BRANDED)["expected_effective"]["body_paragraph_1_run"]
    font = (
        _open(BRANDED).slides[0].placeholders[1].text_frame.paragraphs[1].runs[0].effective_font()
    )
    assert font.size.value_pt == expected["size_centipoints"] / 100.0
    assert _supplied_levels(font.size) == ["master txStyles bodyStyle lvl2"]
    assert font.name.value == expected["font_name"]  # -- Calibri via +mn-lt
    assert _supplied_levels(font.name) == ["theme fontScheme minorFont"]


def test_every_consulted_level_appears_in_provenance_in_walk_order():
    font = (
        _open(BRANDED).slides[0].placeholders[1].text_frame.paragraphs[1].runs[0].effective_font()
    )
    levels = [step.level for step in font.size.provenance]
    assert levels == [
        "run rPr",
        "paragraph defRPr",
        "shape lstStyle lvl2",
        "layout placeholder lstStyle lvl2",
        "master placeholder lstStyle lvl2",
        "master txStyles bodyStyle lvl2",
    ]


# ------------------------------------------------------------- scheme-color resolution


def test_scheme_color_resolves_through_remapped_clrmap():
    expected = _ground_truth(CLRMAP)["expected_resolution"]
    prs = _open(CLRMAP)
    box = next(s for s in prs.slides[0].shapes if s.name == "tx1_text")
    font = box.text_frame.paragraphs[0].runs[0].effective_font()
    assert font.color_rgb.resolved
    assert font.color_rgb.value == expected["text_run_rgb"]
    details = [step.detail for step in font.color_rgb.provenance]
    assert 'schemeClr val="tx1"' in details[0]
    assert 'tx1="lt1"' in details  # -- the clrMap remap step is visible
    assert _supplied_levels(font.color_rgb) == ["theme clrScheme lt1"]


def test_slide_clrmapovr_override_beats_master_clrmap():
    prs = _open(CLRMAP)
    slide = prs.slides[0]
    clrMapOvr = slide._element.get_or_add_clrMapOvr()
    override = etree.SubElement(clrMapOvr, "{%s}overrideClrMapping" % _A)
    master_map = dict(prs.slide_masters[0].element.find(
        "{http://schemas.openxmlformats.org/presentationml/2006/main}clrMap").attrib)
    master_map["tx1"] = "accent3"  # -- 9BBB59 in the default theme
    for key, value in master_map.items():
        override.set(key, value)

    box = next(s for s in slide.shapes if s.name == "tx1_text")
    font = box.text_frame.paragraphs[0].runs[0].effective_font()
    assert font.color_rgb.value == "9BBB59"
    assert any(step.level == "slide clrMapOvr" for step in font.color_rgb.provenance)


def test_direct_srgb_color_supplies_immediately():
    prs = _open(BRANDED)
    run = prs.slides[0].shapes.title.text_frame.paragraphs[0].runs[0]
    run.font.color.rgb = RGBColor(0x12, 0x34, 0xAB)
    font = run.effective_font()
    assert font.color_rgb.value == "1234AB"
    assert _supplied_levels(font.color_rgb) == ["run rPr"]


@pytest.mark.parametrize("color_tag", ["srgbClr", "schemeClr"])
def test_transformed_color_is_unresolved_instead_of_reporting_base_rgb(color_tag):
    prs = _open("self_generated/minimal_clean.pptx")
    box = prs.slides[0].shapes.add_textbox(0, 0, 914400, 914400)
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "transformed"
    rPr = run._r.get_or_add_rPr()
    solidFill = etree.SubElement(rPr, "{%s}solidFill" % _A)
    color = etree.SubElement(solidFill, "{%s}%s" % (_A, color_tag))
    color.set("val", "804020" if color_tag == "srgbClr" else "accent1")
    etree.SubElement(color, "{%s}lumMod" % _A).set("val", "50000")

    effective = run.effective_font().color_rgb
    assert not effective.resolved
    assert effective.value is None
    assert any("unapplied transforms" in step.detail for step in effective.provenance)


# ----------------------------------------------------- non-placeholder + honesty cases


def test_plain_textbox_resolves_via_presentation_default_text_style():
    prs = _open("self_generated/minimal_clean.pptx")
    box = prs.slides[0].shapes.add_textbox(0, 0, 914400, 914400)
    box.text_frame.paragraphs[0].add_run().text = "plain"
    font = box.text_frame.paragraphs[0].runs[0].effective_font()
    assert (font.size.value_pt, font.name.value, font.color_rgb.value) == (
        18.0,
        "Calibri",
        "000000",
    )
    assert _supplied_levels(font.size) == ["presentation defaultTextStyle lvl1"]


def test_autoshape_fontref_supplies_theme_font_and_text_color():
    prs = _open("self_generated/minimal_clean.pptx")
    shape = prs.slides[0].shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, 914400, 914400)
    shape.text = "font ref"
    font = shape.text_frame.paragraphs[0].runs[0].effective_font()
    assert font.name.value == "Calibri"
    assert font.color_rgb.value == "FFFFFF"
    assert "shape fontRef" in [step.level for step in font.name.provenance]


def test_non_latin_run_does_not_report_the_latin_typeface_as_effective():
    prs = _open("self_generated/minimal_clean.pptx")
    box = prs.slides[0].shapes.add_textbox(0, 0, 914400, 914400)
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "日本語"
    run.font.name = "Arial"
    etree.SubElement(run._r.get_or_add_rPr(), "{%s}ea" % _A).set("typeface", "Yu Mincho")

    name = run.effective_font().name
    assert not name.resolved
    assert name.value is None
    assert "non-Latin" in name.provenance[-1].detail


def test_gradient_text_fill_reports_unresolved_not_a_guess():
    prs = _open("self_generated/minimal_clean.pptx")
    box = prs.slides[0].shapes.add_textbox(0, 0, 914400, 914400)
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "gradient"
    rPr = run._r.get_or_add_rPr()
    etree.SubElement(rPr, "{%s}gradFill" % _A)
    font = box.text_frame.paragraphs[0].runs[0].effective_font()
    assert not font.color_rgb.resolved
    assert font.color_rgb.value is None
    assert "gradFill" in font.color_rgb.provenance[-1].detail


def test_phclr_scheme_token_reports_unresolved():
    prs = _open("self_generated/minimal_clean.pptx")
    box = prs.slides[0].shapes.add_textbox(0, 0, 914400, 914400)
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "phClr"
    rPr = run._r.get_or_add_rPr()
    solidFill = etree.SubElement(rPr, "{%s}solidFill" % _A)
    schemeClr = etree.SubElement(solidFill, "{%s}schemeClr" % _A)
    schemeClr.set("val", "phClr")
    font = run.effective_font()
    assert not font.color_rgb.resolved
    assert any("unmappable" in step.detail for step in font.color_rgb.provenance)


def test_table_cell_run_refuses_instead_of_guessing():
    prs = _open(GAUNTLET)
    table_shape = next(s for s in prs.slides[2].shapes if s.name == "gauntlet_table")
    cell_run = table_shape.table.cell(0, 0).text_frame.paragraphs[0].runs[0]
    with pytest.raises(UnsupportedStructureError):
        effective_font(cell_run)


# ---------------------------------------------------------------- read-only + determinism


def test_inspection_never_mutates_the_package():
    prs = _open(BRANDED)
    before = snapshot_parts(prs)
    inspect_text(prs.slides[0])
    for shape in prs.slides[0].shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.effective_font()
    assert snapshot_parts(prs) == before


def test_inspection_is_deterministic_across_runs_and_loads():
    first = json.dumps(inspect_text(_open(BRANDED).slides[0]).to_dict())
    second = json.dumps(inspect_text(_open(BRANDED).slides[0]).to_dict())
    assert first == second


@pytest.mark.parametrize(
    ("golden_name", "fixture_relpath", "slide_index"),
    [
        ("branded_template.inspect.json", BRANDED, 0),
        ("clrmap_remap.inspect.json", CLRMAP, 0),
        ("gauntlet_slide1.inspect.json", GAUNTLET, 0),
    ],
)
def test_inspection_matches_frozen_golden(golden_name, fixture_relpath, slide_index):
    """Byte-identical to the reviewed golden; update ONLY via update_goldens.py + PR review."""
    golden_path = corpus.FIXTURES_DIR.parent / "goldens" / golden_name
    payload = inspect_text(_open(fixture_relpath).slides[slide_index]).to_dict()
    actual = (json.dumps(payload, indent=2, ensure_ascii=False) + "\n").encode("utf-8")
    assert actual == golden_path.read_bytes()  # -- byte-exact, no newline translation


def test_libreoffice_independently_confirms_branded_effective_sizes():
    """LO resolved the same inheritance when it baked run sizes into its re-export."""
    ours = sorted(
        run.font.size.value_pt
        for block in inspect_text(_open(BRANDED).slides[0]).blocks
        for run in block.runs
    )
    lo_prs = Presentation(str(corpus.fixture_path("libreoffice_export/lo_branded_template.pptx")))
    theirs = sorted(
        r.effective_font().size.value_pt
        for s in lo_prs.slides[0].shapes
        if s.has_text_frame
        for p in s.text_frame.paragraphs
        for r in p.runs
    )
    assert ours == theirs == [22.0, 26.0, 36.0]


# --------------------------------------------------------------------------------- anchors


# -------------------------------------------------- visibility completeness


def test_inspect_text_sees_table_cells_as_counted_blind_regions():
    """Table-cell text appears (row-major) as typed blind regions, never silence."""
    inspection = inspect_text(_open("self_generated/tables_in_group.pptx").slides[0])
    by_container = {}
    for block in inspection.blocks:
        by_container.setdefault(block.container, []).append(block)

    cell_blocks = by_container.get("table-cell", [])
    assert [b.text for b in cell_blocks] == [
        "cell r0c0", "cell r0c1", "cell r1c0", "cell r1c1"
    ]
    assert [b.container_detail for b in cell_blocks] == [
        "grouped_table!r0c0", "grouped_table!r0c1",
        "grouped_table!r1c0", "grouped_table!r1c1",
    ]
    assert all(b.blind for b in cell_blocks)
    for block in cell_blocks:
        for run in block.runs:
            assert run.font.size.resolved is False  # -- honest unresolved, not a guess
    assert inspection.blind_region_count == 4

    payload = inspection.to_dict()
    assert payload["version"] == 2
    assert payload["blind_region_count"] == 4


def test_inspect_text_sees_grouped_shape_text_with_group_paths():
    """Text inside groups appears, recursively, with its group path."""
    inspection = inspect_text(_open("self_generated/nested_groups.pptx").slides[0])
    by_text = {b.text: b for b in inspection.blocks}
    assert by_text["Level zero"].container == "shape"
    assert by_text["Level zero"].container_detail is None
    assert by_text["Level one"].container == "group"
    assert by_text["Level one"].container_detail == "group_level1"
    assert by_text["Level two"].container_detail == "group_level1/group_level2"
    assert by_text["Level three"].container_detail == (
        "group_level1/group_level2/group_level3"
    )
    assert not any(b.blind for b in inspection.blocks)
    # -- grouped runs resolve through the normal (non-placeholder) chain
    assert by_text["Level three"].runs[0].font.size.resolved is True


def test_inspect_text_sees_in_group_textbox_beside_grouped_table():
    inspection = inspect_text(_open("self_generated/tables_in_group.pptx").slides[0])
    by_text = {b.text: b for b in inspection.blocks}
    assert by_text["Top-level text"].container == "shape"
    assert by_text["In-group text"].container == "group"
    assert by_text["In-group text"].container_detail == "outer_group"


def test_inspect_text_block_order_is_depth_first_document_order():
    inspection = inspect_text(_open("self_generated/tables_in_group.pptx").slides[0])
    texts = [b.text for b in inspection.blocks]
    assert texts == [
        "Top-level text", "In-group text",
        "cell r0c0", "cell r0c1", "cell r1c0", "cell r1c1",
    ]
    assert [b.anchor.block_index for b in inspection.blocks] == list(range(len(texts)))


@pytest.mark.parametrize("depth", [17, 64, 248])
def test_deeply_nested_groups_inspect_without_refusing(depth):
    """A retired guard used to refuse the whole slide past depth 16, hiding ordinary shapes
    alongside the nest. Depth 248 is the deepest a deck can be and still parse: libxml2 caps
    document nesting at 256 elements, so no loadable deck exceeds it."""
    prs = _open("self_generated/minimal_clean.pptx")
    slide = prs.slides[0]
    shallow_before = len(inspect_text(slide).blocks)

    group = slide.shapes.add_group_shape()
    for _ in range(depth - 1):
        group = group.shapes.add_group_shape()
    box = group.shapes.add_textbox(0, 0, 914400, 914400)
    box.text_frame.paragraphs[0].add_run().text = "deep text"

    blocks = inspect_text(slide).blocks
    texts = [b.text for b in blocks]
    # -- the deep block is reported, and every pre-existing shallow block survives with it
    assert "deep text" in texts
    assert len(blocks) == shallow_before + 1
    deep = next(b for b in blocks if b.text == "deep text")
    assert deep.container == "group"
    assert deep.blind is False


def test_deeply_nested_group_deck_round_trips_through_a_file():
    """The nest must survive save -> load, not just in-memory inspection."""
    import io

    prs = _open("self_generated/minimal_clean.pptx")
    group = prs.slides[0].shapes.add_group_shape()
    for _ in range(247):
        group = group.shapes.add_group_shape()
    box = group.shapes.add_textbox(0, 0, 914400, 914400)
    box.text_frame.paragraphs[0].add_run().text = "deep text"

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    reloaded = Presentation(buf)
    assert "deep text" in [b.text for b in inspect_text(reloaded.slides[0]).blocks]


def test_effective_font_resolves_runs_inside_groups():
    prs = _open("self_generated/nested_groups.pptx")
    group = next(s for s in prs.slides[0].shapes if s.name == "group_level1")
    for name in ("group_level2", "group_level3"):
        group = next(s for s in group.shapes if s.name == name)
    box = next(s for s in group.shapes if s.name == "level3_box")
    info = box.text_frame.paragraphs[0].runs[0].effective_font()
    assert info.size.resolved is True


def test_content_hash_is_pinned_sha256_nfc_prefix():
    assert content_hash("Branded Title") == (
        hashlib.sha256("Branded Title".encode("utf-8")).hexdigest()[:8]
    )


def test_content_hash_applies_nfc_normalization():
    composed = "café"  # -- é as one code point
    decomposed = "café"  # -- e + combining acute
    assert composed != decomposed
    assert content_hash(composed) == content_hash(decomposed)


def test_out_of_schema_indent_level_refuses_instead_of_crashing():
    prs = _open(BRANDED)
    run = prs.slides[0].shapes.title.text_frame.paragraphs[0].runs[0]
    pPr = run._r.getparent().find("{%s}pPr" % _A)
    if pPr is None:
        pPr = etree.SubElement(run._r.getparent(), "{%s}pPr" % _A)
        run._r.getparent().insert(0, pPr)
    pPr.set("lvl", "9")  # -- outside ST_TextIndentLevelType's 0..8
    with pytest.raises(UnsupportedStructureError, match="lvl=9"):
        run.effective_font()


def test_alternate_content_is_a_typed_counted_blind_region_not_silence():
    """Regression (review): mc:AlternateContent used to be invisible to inspect_text and
    inspect_deck — a fail-silent hole in the visibility-complete contract."""
    from pptx.inspect import inspect_deck
    from tests.paper.test_edit_text import _wrap_first_textbox_in_alternate_content

    prs = _wrap_first_textbox_in_alternate_content(
        _open("self_generated/minimal_clean.pptx")
    )
    inspection = inspect_text(prs.slides[0])
    ac_blocks = [b for b in inspection.blocks if b.container == "alternate-content"]
    assert len(ac_blocks) == 1
    assert ac_blocks[0].blind is True
    assert inspection.blind_region_count == 1
    assert inspection.to_dict()["blind_region_count"] == 1

    manifest = inspect_deck(prs)
    assert manifest.slides[0].alternate_content_count == 1
    assert manifest.to_dict()["slides"][0]["alternate_content_count"] == 1


def test_boolean_effective_values_serialize_as_json_booleans():
    """Regression (review): bool is an int subclass; the payload used to emit 0/1."""
    payload = (
        _open(BRANDED).slides[0].placeholders[1].text_frame.paragraphs[0].runs[0]
        .effective_font()
        .to_dict()
    )
    assert payload["bold"]["value"] is False
    assert payload["italic"]["value"] is False


# --------------------------------------------------------- walk extensions


def test_bold_italic_underline_resolve_with_schema_defaults():
    prs = _open(BRANDED)
    run = prs.slides[0].placeholders[1].text_frame.paragraphs[0].runs[0]
    font = run.effective_font()
    assert font.bold.value is False
    assert font.bold.resolved is True
    assert font.bold.provenance[-1].level == "schema default"
    assert font.italic.value is False
    assert font.italic.resolved is True
    assert font.underline.value == "none"
    assert font.underline.resolved is True

    run.font.bold = True
    run.font.underline = True  # -- upstream writes u="sng"
    font = run.effective_font()
    assert font.bold.value is True
    assert font.bold.provenance[-1].level == "run rPr"
    assert font.underline.value == "sng"


def test_effective_paragraph_format_resolves_alignment_and_spacing():
    from pptx.enum.text import PP_ALIGN

    prs = _open(BRANDED)
    paragraph = prs.slides[0].placeholders[1].text_frame.paragraphs[0]
    fmt = effective_paragraph_format(paragraph)
    assert fmt.alignment.value == "l"  # -- master bodyStyle supplies algn="l"
    assert any(s.supplied and "bodyStyle" in s.level for s in fmt.alignment.provenance)
    assert fmt.line_spacing.value == 1.0  # -- rendering default, explicitly provenanced
    assert fmt.line_spacing.provenance[-1].level == "rendering default"

    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.line_spacing = 1.5
    fmt = effective_paragraph_format(paragraph)
    assert fmt.alignment.value == "ctr"
    assert fmt.alignment.provenance[0].supplied is True  # -- paragraph pPr wins
    assert fmt.line_spacing.value == 1.5

    payload = fmt.to_dict()
    assert payload["schema"] == "paper-effective-paragraph-format"
    assert payload["version"] == 3  # -- v3 added bullet_font/bullet_size


# ----------------------------------------------------------- inherited bullet resolution


def _bullet(paragraph):
    return effective_paragraph_format(paragraph).bullet


def _graft_lst_style_bullet(shape, lvl_tag, bullet_tag, **attrs):
    """Add `a:lstStyle/<lvl_tag>/<bullet_tag>` to `shape`'s `a:txBody`, in memory only.

    `get_or_add_lstStyle` places `a:lstStyle` at its schema position in `CT_TextBody`
    (after `a:bodyPr`, before the first `a:p`); appending it would invalidate the fragment.
    """
    lst_style = shape.text_frame._txBody.get_or_add_lstStyle()
    level_properties = etree.SubElement(lst_style, "{%s}%s" % (_A, lvl_tag))
    bullet = etree.SubElement(level_properties, "{%s}%s" % (_A, bullet_tag))
    for name, value in attrs.items():
        bullet.set(name, value)
    return bullet


def _plain_textbox(prs):
    """A non-placeholder shape: its chain takes the `presentation defaultTextStyle` branch."""
    box = prs.slides[0].shapes.add_textbox(0, 0, Emu(914400 * 4), Emu(914400))
    box.text_frame.paragraphs[0].add_run().text = "no bullet anywhere"
    return box


def test_bullet_resolves_the_inherited_character_bullet_from_the_master():
    """The regression this feature exists for: `BulletFormat.type` reads |None| here."""
    paragraph = _open(BRANDED).slides[0].placeholders[1].text_frame.paragraphs[0]
    assert paragraph.bullet.type is None  # -- local state says nothing

    bullet = _bullet(paragraph)

    assert bullet.type == "character"
    assert bullet.char == "•"
    assert bullet.resolved is True
    assert bullet.number_scheme is None
    assert bullet.start_at is None
    assert _supplied_levels(bullet) == ["master txStyles bodyStyle lvl1"]


def test_bullet_resolves_per_indent_level_not_always_level_one():
    """The lvl=1 paragraph must pick up bodyStyle lvl2's en dash, not lvl1's bullet."""
    paragraph = _open(BRANDED).slides[0].placeholders[1].text_frame.paragraphs[1]

    bullet = _bullet(paragraph)

    assert bullet.type == "character"
    assert bullet.char == "–"
    assert _supplied_levels(bullet) == ["master txStyles bodyStyle lvl2"]


def test_bullet_reports_master_bunone_as_a_positive_terminal():
    """`buNone` is an answer, not an absence: the master supplies it, not the fallback."""
    paragraph = _open(BRANDED).slides[0].shapes.title.text_frame.paragraphs[0]

    bullet = _bullet(paragraph)

    assert bullet.type == "none"
    assert bullet.resolved is True
    assert bullet.char is None
    # -- a resolver that walked past the buNone would also say "none", but would land on the
    # -- synthetic "rendering default" step instead; only the supplying rung tells them apart
    assert _supplied_levels(bullet) == ["master txStyles titleStyle lvl1"]


def test_bullet_resolves_a_local_numbered_override_at_the_first_rung():
    paragraph = _open(BRANDED).slides[0].placeholders[1].text_frame.paragraphs[0]
    paragraph.bullet.set_numbered("arabicPeriod", start_at=3)

    bullet = _bullet(paragraph)

    assert bullet.type == "numbered"
    assert bullet.number_scheme == "arabicPeriod"
    assert bullet.start_at == 3
    assert bullet.char is None
    assert _supplied_levels(bullet) == ["paragraph pPr"]


def test_bullet_resolves_to_none_when_the_whole_chain_is_exhausted():
    """Exhaustion resolves — there is no ECMA default bullet, so provenance says so."""
    prs = _open(BRANDED)
    paragraph = _plain_textbox(prs).text_frame.paragraphs[0]

    bullet = _bullet(paragraph)

    assert bullet.type == "none"
    assert bullet.resolved is True
    assert _supplied_levels(bullet) == ["rendering default"]
    assert bullet.provenance[-1].part is None
    assert bullet.provenance[-1].supplied is True


def test_bullet_shape_lst_style_bunone_shadows_the_master_buchar():
    prs = _open(BRANDED)
    body = prs.slides[0].placeholders[1]
    _graft_lst_style_bullet(body, "lvl1pPr", "buNone")

    bullet = _bullet(body.text_frame.paragraphs[0])

    assert bullet.type == "none"
    assert bullet.resolved is True
    assert _supplied_levels(bullet) == ["shape lstStyle lvl1"]


def test_bullet_sees_an_inherited_bublip_outside_the_choice_group():
    """`a:buBlip` is not an `eg_bullet` member; reading only the descriptor would walk past."""
    prs = _open(BRANDED)
    body = prs.slides[0].placeholders[1]
    buBlip = _graft_lst_style_bullet(body, "lvl2pPr", "buBlip")
    etree.SubElement(buBlip, "{%s}blip" % _A)

    bullet = _bullet(body.text_frame.paragraphs[1])  # -- the lvl=1 paragraph

    assert bullet.type == "picture"
    assert bullet.resolved is True
    assert bullet.char is None
    assert bullet.number_scheme is None
    assert bullet.start_at is None
    assert _supplied_levels(bullet) == ["shape lstStyle lvl2"]


@pytest.mark.parametrize(
    ("bullet_tag", "attrs"),
    [
        ("buChar", {}),  # -- schema-required @char missing
        ("buAutoNum", {}),  # -- schema-required @type missing
        ("buAutoNum", {"type": "arabicPeriod", "startAt": "many"}),  # -- @startAt not an int
    ],
)
def test_bullet_reports_a_malformed_member_as_unresolved(bullet_tag, attrs):
    """Honest over plausible: no guessed value, and no walking on to the master's bullet.

    Reading these through the generated `RequiredAttribute` descriptors would raise
    `InvalidXmlError` instead, turning the honesty rule into a crash on malformed input.
    """
    prs = _open(BRANDED)
    body = prs.slides[0].placeholders[1]
    _graft_lst_style_bullet(body, "lvl1pPr", bullet_tag, **attrs)

    bullet = _bullet(body.text_frame.paragraphs[0])

    assert bullet.resolved is False
    assert bullet.type is None
    assert (bullet.char, bullet.number_scheme, bullet.start_at) == (None, None, None)
    assert _supplied_levels(bullet) == []
    # -- provenance is intact up to and including the offending rung ...
    offender = bullet.provenance[-1]
    assert offender.supplied is False
    assert bullet_tag in offender.detail
    # -- ... and the walk stopped there rather than reporting the master's bullet
    assert [step.level for step in bullet.provenance] == ["paragraph pPr", "shape lstStyle lvl1"]


def test_bullet_resolves_an_inherited_numbered_bullet_and_defaults_start_at_to_one():
    """`@startAt` is optional on `a:buAutoNum`; absent means the sequence starts at 1."""
    prs = _open(BRANDED)
    body = prs.slides[0].placeholders[1]
    _graft_lst_style_bullet(body, "lvl1pPr", "buAutoNum", type="romanUcParenR")

    bullet = _bullet(body.text_frame.paragraphs[0])

    assert bullet.type == "numbered"
    assert bullet.number_scheme == "romanUcParenR"
    assert bullet.start_at == 1
    assert bullet.char is None
    assert _supplied_levels(bullet) == ["shape lstStyle lvl1"]


def test_bullet_provenance_lists_every_consulted_rung_in_walk_order():
    prs = _open(BRANDED)
    placeholder_bullet = _bullet(prs.slides[0].placeholders[1].text_frame.paragraphs[1])
    assert [step.level for step in placeholder_bullet.provenance] == [
        "paragraph pPr",
        "shape lstStyle lvl2",
        "layout placeholder lstStyle lvl2",
        "master placeholder lstStyle lvl2",
        "master txStyles bodyStyle lvl2",
    ]
    assert _supplied_levels(placeholder_bullet) == ["master txStyles bodyStyle lvl2"]
    # -- a silent rung says which silence it is: rung absent, or rung present with no bullet
    assert [step.detail for step in placeholder_bullet.provenance[:-1]] == [
        "no bullet here",  # -- this paragraph does have an a:pPr; it carries only lvl="1"
        "level not present",
        "level not present",
        "level not present",
    ]

    textbox_bullet = _bullet(_plain_textbox(prs).text_frame.paragraphs[0])
    assert [step.level for step in textbox_bullet.provenance] == [
        "paragraph pPr",
        "shape lstStyle lvl1",
        "presentation defaultTextStyle lvl1",
        "rendering default",
    ]
    assert _supplied_levels(textbox_bullet) == ["rendering default"]


def test_bullet_resolution_never_mutates_the_package():
    """A `get_or_add_*` read would materialize an empty `a:pPr` and change the saved bytes."""
    prs = _open(BRANDED)
    paragraphs = [
        paragraph
        for shape in prs.slides[0].shapes  # -- title (buNone path) and body (buChar path)
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
    ]
    before = snapshot_parts(prs)

    resolved = [_bullet(paragraph) for paragraph in paragraphs]

    assert [bullet.type for bullet in resolved] == ["none", "character", "character"]
    assert snapshot_parts(prs) == before


def test_bullet_payload_serializes_in_the_documented_shape():
    prs = _open(BRANDED)
    body = prs.slides[0].placeholders[1]

    payload = effective_paragraph_format(body.text_frame.paragraphs[0]).to_dict()
    assert payload["schema"] == "paper-effective-paragraph-format"
    assert payload["version"] == 3
    bullet = payload["bullet"]
    assert sorted(bullet) == ["char", "number_scheme", "provenance", "resolved", "start_at", "type"]
    assert bullet["type"] == "character"
    assert bullet["char"] == "•"
    assert bullet["number_scheme"] is None
    assert bullet["start_at"] is None
    assert bullet["resolved"] is True
    assert bullet["provenance"][-1] == {
        "level": "master txStyles bodyStyle lvl1",
        "part": "/ppt/slideMasters/slideMaster1.xml",
        "detail": 'buChar char="•"',
        "supplied": True,
    }

    body.text_frame.paragraphs[0].bullet.set_numbered("arabicPeriod", start_at=3)
    numbered = effective_paragraph_format(body.text_frame.paragraphs[0]).to_dict()["bullet"]
    assert numbered["type"] == "numbered"
    assert numbered["char"] is None  # -- char is populated for character bullets only
    assert numbered["number_scheme"] == "arabicPeriod"
    assert numbered["start_at"] == 3

    title = prs.slides[0].shapes.title
    none_bullet = effective_paragraph_format(title.text_frame.paragraphs[0]).to_dict()["bullet"]
    assert none_bullet["type"] == "none"
    assert (none_bullet["char"], none_bullet["number_scheme"], none_bullet["start_at"]) == (
        None,
        None,
        None,
    )


# ------------------------------------------------- bullet typeface and size resolution

WINGDING = ""  # -- private use: a filled square in Wingdings, undrawable elsewhere


def test_bullet_font_and_size_round_trip_through_the_write_api():
    """The defect this feature exists for.

    A private-use glyph is meaningless without its typeface, so a caller re-applying
    everything the payload reports must land on the same `buFont` it started from. Before
    `bullet_font` existed there was nothing to pass and the copy rendered as a blank box.
    """
    prs = _open(BRANDED)
    body = prs.slides[0].placeholders[1]
    source = body.text_frame.paragraphs[0]
    source.bullet.set_character(WINGDING, font_name="Wingdings", size_percent=0.75)

    fmt = effective_paragraph_format(source)
    assert fmt.bullet.char == WINGDING
    assert fmt.bullet_font.value == "Wingdings"
    assert fmt.bullet_size.value == 0.75

    target = body.text_frame.add_paragraph()
    target.text = "re-applied from the payload alone"
    target.bullet.set_character(
        fmt.bullet.char, font_name=fmt.bullet_font.value, size_percent=fmt.bullet_size.value
    )

    reopened = save_reopen(prs)
    shape = next(s for s in reopened.slides[0].shapes if s.shape_id == body.shape_id)
    copied = effective_paragraph_format(shape.text_frame.paragraphs[-1])
    assert copied.bullet.char == WINGDING
    assert copied.bullet_font.value == "Wingdings"  # -- the typeface survived the round trip
    assert copied.bullet_size.value == 0.75


def test_bullet_font_resolves_independently_of_the_bullet_kind():
    """Three separate choice groups, three separate walks, three provenance chains."""
    prs = _open(BRANDED)
    body = prs.slides[0].placeholders[1]
    _graft_lst_style_bullet(body, "lvl1pPr", "buFont", typeface="Wingdings")

    fmt = effective_paragraph_format(body.text_frame.paragraphs[0])
    # -- the kind still comes from the master, the typeface from the shape's own lstStyle
    assert fmt.bullet_font.value == "Wingdings"
    assert _supplied_levels(fmt.bullet_font) == ["shape lstStyle lvl1"]
    assert _supplied_levels(fmt.bullet) == ["master txStyles bodyStyle lvl1"]


def test_bullet_font_follows_text_explicitly_and_on_exhaustion():
    """`buFontTx` is an answer attributed to its rung; exhaustion reaches it as a default."""
    prs = _open(BRANDED)
    body = prs.slides[0].placeholders[1]
    _graft_lst_style_bullet(body, "lvl1pPr", "buFontTx")
    explicit = effective_paragraph_format(body.text_frame.paragraphs[0]).bullet_font
    assert explicit.value == BULLET_FOLLOWS_TEXT
    assert explicit.resolved is True
    assert _supplied_levels(explicit) == ["shape lstStyle lvl1"]

    textbox = _plain_textbox(_open(BRANDED))
    exhausted = effective_paragraph_format(textbox.text_frame.paragraphs[0]).bullet_font
    assert exhausted.value == BULLET_FOLLOWS_TEXT
    assert exhausted.resolved is True
    assert _supplied_levels(exhausted) == ["rendering default"]


def test_bullet_size_discriminates_its_three_members():
    """`buSzPct` is a fraction, `buSzPts` an absolute length, `buSzTx` defers to the text."""
    pct = _open(BRANDED)
    _graft_lst_style_bullet(pct.slides[0].placeholders[1], "lvl1pPr", "buSzPct", val="45%")
    resolved = effective_paragraph_format(
        pct.slides[0].placeholders[1].text_frame.paragraphs[0]
    ).bullet_size
    assert resolved.value == 0.45
    assert resolved.value_pt is None

    pts = _open(BRANDED)
    _graft_lst_style_bullet(pts.slides[0].placeholders[1], "lvl1pPr", "buSzPts", val="1400")
    resolved = effective_paragraph_format(
        pts.slides[0].placeholders[1].text_frame.paragraphs[0]
    ).bullet_size
    assert resolved.value_pt == 14.0  # -- centipoints, as spcPts is reported

    tx = _open(BRANDED)
    _graft_lst_style_bullet(tx.slides[0].placeholders[1], "lvl1pPr", "buSzTx")
    resolved = effective_paragraph_format(
        tx.slides[0].placeholders[1].text_frame.paragraphs[0]
    ).bullet_size
    assert resolved.value == BULLET_FOLLOWS_TEXT


def test_bullet_font_resolves_a_theme_token_through_the_font_scheme():
    """A `buFont` may name `+mj-lt`; report the typeface it lands on, not the token."""
    prs = _open(BRANDED)
    body = prs.slides[0].placeholders[1]
    _graft_lst_style_bullet(body, "lvl1pPr", "buFont", typeface="+mj-lt")

    resolved = effective_paragraph_format(body.text_frame.paragraphs[0]).bullet_font
    assert resolved.resolved is True
    assert resolved.value == _ground_truth(BRANDED)["theme_major_latin"]
    assert any("theme reference" in step.detail for step in resolved.provenance)
    assert _supplied_levels(resolved) == ["theme fontScheme majorFont"]


@pytest.mark.parametrize(
    ("tag", "attrs", "detail"),
    [
        ("buFont", {}, "buFont with no typeface"),
        ("buSzPct", {"val": "not-a-number"}, "buSzPct with unreadable"),
        ("buSzPts", {"val": "oops"}, "buSzPts with unreadable"),
    ],
)
def test_bullet_modifier_missing_its_required_attribute_reports_unresolved(tag, attrs, detail):
    """Honest over plausible: report, do not guess, and do not raise |InvalidXmlError|."""
    prs = _open(BRANDED)
    body = prs.slides[0].placeholders[1]
    _graft_lst_style_bullet(body, "lvl1pPr", tag, **attrs)

    fmt = effective_paragraph_format(body.text_frame.paragraphs[0])
    facet = fmt.bullet_font if tag == "buFont" else fmt.bullet_size
    assert facet.resolved is False
    assert facet.value is None
    assert _supplied_levels(facet) == []
    assert detail in facet.provenance[-1].detail
    assert facet.provenance[-1].level == "shape lstStyle lvl1"


def test_bullet_modifier_resolution_never_mutates_the_package():
    prs = _open(BRANDED)
    paragraphs = [
        paragraph
        for shape in prs.slides[0].shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
    ]
    before = snapshot_parts(prs)

    resolved = [effective_paragraph_format(paragraph) for paragraph in paragraphs]

    # -- the title's chain runs through titleStyle, which carries buNone and no buFont
    assert [f.bullet_font.value for f in resolved] == [BULLET_FOLLOWS_TEXT, "Arial", "Arial"]
    assert snapshot_parts(prs) == before


def test_bullet_modifier_payload_serializes_beside_the_bullet():
    prs = _open(BRANDED)
    payload = effective_paragraph_format(
        prs.slides[0].placeholders[1].text_frame.paragraphs[0]
    ).to_dict()
    assert payload["version"] == 3
    assert sorted(payload) == [
        "alignment",
        "bullet",
        "bullet_font",
        "bullet_size",
        "line_spacing",
        "schema",
        "version",
    ]
    for key in ("bullet_font", "bullet_size"):
        assert sorted(payload[key]) == ["provenance", "resolved", "value", "value_pt"]
    assert payload["bullet_font"]["value"] == "Arial"  # -- master bodyStyle supplies buFont
    assert payload["bullet_font"]["provenance"][-1]["detail"] == 'buFont typeface="Arial"'


# ------------------------------------------------------------------ shape-format resolution


def test_effective_shape_format_resolves_explicit_fill_through_clrmap():
    """A probe case: the rectangle behind the text, resolved like the text."""
    from pptx.inspect import effective_shape_format

    prs = _open(CLRMAP)
    rect = prs.slides[0].shapes.shape_by_name("accent1_box")
    fmt = effective_shape_format(rect)
    assert fmt.fill_rgb.value == "C0504D"  # -- accent1 -> clrMap -> theme accent2
    assert fmt.fill_rgb.resolved is True
    assert any("clrMap" in s.detail or "clrMap" in s.level for s in fmt.fill_rgb.provenance)

    # -- line color comes only from the style lnRef: honestly unresolved, reference
    # -- color carried in provenance
    assert fmt.line_rgb.resolved is False
    assert any("C0504D" in s.detail for s in fmt.line_rgb.provenance)

    payload = fmt.to_dict()
    assert payload["schema"] == "paper-effective-shape-format"
    assert payload["version"] == 1


def test_effective_shape_format_reports_nofill_and_absent_fill_honestly():
    from pptx.inspect import effective_shape_format

    prs = _open(CLRMAP)
    rect = prs.slides[0].shapes.shape_by_name("accent1_box")
    rect.fill.background()  # -- a:noFill
    fmt = effective_shape_format(rect)
    assert fmt.fill_rgb.value == "none"
    assert fmt.fill_rgb.resolved is True

    # -- upstream textboxes carry an explicit a:noFill too
    box = prs.slides[0].shapes.shape_by_name("tx1_text")
    assert effective_shape_format(box).fill_rgb.value == "none"

    # -- a placeholder with an empty spPr and no p:style is honestly unresolved
    branded = _open(BRANDED)
    title = branded.slides[0].shapes.title
    fmt = effective_shape_format(title)
    assert fmt.fill_rgb.resolved is False
    assert fmt.fill_rgb.value is None


def test_effective_font_payload_carries_pinned_schema_keys():
    payload = (
        _open(BRANDED).slides[0].shapes.title.text_frame.paragraphs[0].runs[0]
        .effective_font()
        .to_dict()
    )
    assert payload["schema"] == "paper-effective-font"
    assert payload["version"] == 2  # -- v2: bold/italic/underline added


def test_content_hash_treats_whitespace_as_content():
    assert content_hash("Trailing space ") != content_hash("Trailing space")


def test_blocks_carry_stable_anchors_in_sptree_order():
    inspection = inspect_text(_open(BRANDED).slides[0])
    assert [block.anchor.block_index for block in inspection.blocks] == list(
        range(len(inspection.blocks))
    )
    title_block = inspection.blocks[0]
    assert title_block.anchor.part == "/ppt/slides/slide1.xml"
    assert title_block.anchor.content_hash == content_hash("Branded Title")
    assert title_block.placeholder_type == "TITLE"
