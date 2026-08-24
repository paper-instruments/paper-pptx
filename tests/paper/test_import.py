"""Contract tests: slide import and deck merge.

The plan's required list: cross-contamination (source byte-identical after edits to the
import), dedupe (three slides from one source -> one transplanted master), every mode
against the two-template corpus, relint + section scan + LO smoke on all outputs, and
determinism goldens on the import report.
"""

from __future__ import annotations

import copy
import io
import json

import pytest

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.errors import (
    AmbiguousTargetError,
    PaperRefusal,
    RelationshipPolicyError,
    TargetNotFoundError,
    UnsupportedStructureError,
)
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from . import corpus
from .contract import (
    assert_changed_parts,
    assert_refusal_atomic,
    save_reopen,
    save_to_bytes,
    zip_member_map,
)
from .idlists import dangling_section_slide_ids, duplicate_section_slide_ids
from .lo import lo_load_smoke
from .relint import dangling_relationship_targets, missing_relationship_references


def test_import_refuses_a_source_relationship_target_owned_by_another_package():
    dest = Presentation()
    dest.slides.add_slide(dest.slide_layouts[6])
    source = _open("self_generated/minimal_clean.pptx")
    foreign = _open("self_generated/gauntlet.pptx")
    foreign_picture = foreign.slides[1].shapes.picture_by_name("gauntlet_img_1")
    foreign_image = foreign_picture.part.related_part(foreign_picture._pic.blip_rId)
    source.slides[0].part.relate_to(foreign_image, RT.IMAGE)
    before = zip_member_map(save_to_bytes(dest))

    with pytest.raises(RelationshipPolicyError, match="owned by another package"):
        dest.import_slide(source, 0, mode="adopt_theme")

    assert zip_member_map(save_to_bytes(dest)) == before


ALPHA = "self_generated/template_alpha.pptx"
BETA = "self_generated/template_beta.pptx"
LO_ALPHA = "libreoffice_export/lo_template_alpha.pptx"
SECTIONS = "self_generated/sections.pptx"
GAUNTLET = "self_generated/gauntlet.pptx"
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
SECTION_IDS = {
    "Intro": "{11111111-1111-4111-8111-111111111111}",
    "Body": "{22222222-2222-4222-8222-222222222222}",
    "Close": "{33333333-3333-4333-8333-333333333333}",
}


def _open(relpath):
    return Presentation(str(corpus.fixture_path(relpath)))


def _sections(prs):
    return prs._element.findall(".//{%s}sectionLst/{%s}section" % (P14, P14))


def _section_by_id(prs, section_id):
    return next(section for section in _sections(prs) if section.get("id") == section_id)


def _section_slide_ids(section):
    return [int(entry.get("id")) for entry in section.findall(".//{%s}sldId" % P14)]


def _set_content_slots(layout, slots):
    placeholders = sorted(
        (ph for ph in layout.placeholders if ph.element.ph_type == PP_PLACEHOLDER.OBJECT),
        key=lambda ph: ph.element.ph_idx,
    )
    assert len(placeholders) == len(slots)
    for placeholder, (ph_type, idx) in zip(placeholders, slots):
        placeholder.element.ph.type = ph_type
        placeholder.element.ph.idx = idx


def _duplicate_content_slot(layout, idx):
    source = next(
        ph for ph in layout.placeholders if ph.element.ph_type == PP_PLACEHOLDER.OBJECT
    )
    duplicate = copy.deepcopy(source._element)
    duplicate.nvSpPr.cNvPr.set("id", "99")
    duplicate.nvSpPr.cNvPr.set("name", "Additional Content Placeholder")
    duplicate.ph.idx = idx
    source._element.addnext(duplicate)


def _assert_clean(saved_bytes):
    zip_map = zip_member_map(saved_bytes)
    assert dangling_relationship_targets(zip_map) == []
    assert missing_relationship_references(zip_map) == []
    assert dangling_section_slide_ids(zip_map) == []
    assert duplicate_section_slide_ids(zip_map) == []


def _assert_layout_candidates_reported(error, candidates):
    message = str(error)
    assert "target_layout" in message
    for layout in candidates:
        identity = "name=%r, type=%r, part=%s, master=%s" % (
            layout.name,
            layout._element.get("type"),
            layout.part.partname,
            layout.slide_master.part.partname,
        )
        assert identity in message


def _assert_explicit_layout_survives_reopen(dest, source, slide_idx, mode, target_layout):
    target_partname = str(target_layout.part.partname)
    report = dest.import_slide(source, slide_idx, mode=mode, target_layout=target_layout)
    assert report.layout_binding_method == "explicit"
    reopened = save_reopen(dest)
    imported = reopened.slides[-1]
    assert str(imported.slide_layout.part.partname) == target_partname
    return imported


# --------------------------------------------------------------------------- mode contracts


def test_keep_appearance_transplants_chain_with_zero_shifts():
    """The keep-appearance invariant: identical chain, identical resolution - no shifts."""
    dest = _open(ALPHA)
    source = _open(BETA)
    report = dest.import_slide(source, 0, mode="keep_appearance")
    assert report.run_shifts == ()
    assert report.layout_binding_method == "transplant"

    saved = save_to_bytes(dest)
    _assert_clean(saved)
    reopened = Presentation(io.BytesIO(saved))
    assert len(reopened.slide_masters) == 2
    imported = reopened.slides[3]
    assert imported.slide_layout.slide_master is not reopened.slide_masters[0]
    # -- the beta theme travelled: title resolves to beta's Courier New in the dest
    title_font = imported.shapes.title.text_frame.paragraphs[0].runs[0].effective_font()
    assert title_font.name.value == "Courier New"


def test_keep_appearance_dedupes_master_across_three_imports():
    """Three imports from one source: ONE new master and ONE new theme, reused twice."""
    dest = _open(ALPHA)
    source = _open(BETA)
    reports = [dest.import_slide(source, i, mode="keep_appearance") for i in range(3)]
    assert reports[0].parts_reused == ()
    for later in reports[1:]:
        assert "/ppt/slideMasters/slideMaster2.xml" in later.parts_reused

    saved = save_to_bytes(dest)
    zip_map = zip_member_map(saved)
    masters = [
        m for m in zip_map if m.startswith("ppt/slideMasters/slideMaster") and m.endswith(".xml")
    ]
    themes = [m for m in zip_map if m.startswith("ppt/theme/") and m.endswith(".xml")]
    assert sorted(masters) == [
        "ppt/slideMasters/slideMaster1.xml",
        "ppt/slideMasters/slideMaster2.xml",
    ]
    assert sorted(themes) == ["ppt/theme/theme1.xml", "ppt/theme/theme2.xml"]
    _assert_clean(saved)
    # -- the transplanted master accumulated all three layouts
    reopened = Presentation(io.BytesIO(saved))
    assert len(reopened.slide_masters[1].slide_layouts) == 3


def test_adopt_theme_takes_house_style_and_reports_shifts():
    dest = _open(ALPHA)
    report = dest.import_slide(_open(BETA), 0, mode="adopt_theme")
    assert report.layout_binding_method == "name-match"  # -- "Title Slide" collides
    assert report.layout_binding == "/ppt/slideLayouts/slideLayout1.xml"
    shifted_names = {
        (s.before["name"]["value"], s.after["name"]["value"]) for s in report.run_shifts
    }
    assert ("Courier New", "Georgia") in shifted_names  # -- title: beta major -> alpha major
    assert ("Times New Roman", "Verdana") in shifted_names

    saved = save_to_bytes(dest)
    _assert_clean(saved)
    reopened = Presentation(io.BytesIO(saved))
    assert len(reopened.slide_masters) == 1  # -- nothing transplanted
    imported = reopened.slides[3]
    title_font = imported.shapes.title.text_frame.paragraphs[0].runs[0].effective_font()
    assert title_font.name.value == "Georgia"  # -- the house look


def test_adopt_theme_report_records_complete_automatic_placeholder_mapping():
    dest = _open(ALPHA)

    report = dest.import_slide(_open(BETA), 0, mode="adopt_theme")

    assert report.placeholder_map_used == ((0, 0), (1, 1))
    assert report.to_dict()["placeholder_map_used"] == [
        {"source_idx": 0, "target_idx": 0},
        {"source_idx": 1, "target_idx": 1},
    ]


@pytest.mark.parametrize(
    ("target_type", "tier"),
    [
        (PP_PLACEHOLDER.OBJECT, "same-type fallback"),
        (PP_PLACEHOLDER.BODY, "compatible-family fallback"),
    ],
)
def test_adopt_theme_placeholder_ambiguity_refuses_atomically(target_type, tier):
    dest = _open(ALPHA)
    source = _open(BETA)
    source_before = save_to_bytes(source)
    target = dest.slide_layouts[3]
    _set_content_slots(target, [(target_type, 3), (target_type, 2)])
    before = save_to_bytes(dest)

    error = assert_refusal_atomic(
        dest,
        lambda prs: prs.import_slide(
            source, 1, mode="adopt_theme", target_layout=prs.slide_layouts[3]
        ),
        AmbiguousTargetError,
    )

    message = str(error)
    assert "type OBJECT, idx 1" in message
    assert tier in message
    assert "idx 2" in message
    assert "idx 3" in message
    assert "placeholder_map" in message
    assert_changed_parts(before, save_to_bytes(dest))
    assert zip_member_map(save_to_bytes(source)) == zip_member_map(source_before)


def test_adopt_theme_partial_placeholder_map_disambiguates_and_persists():
    dest = _open(ALPHA)
    source = _open(BETA)
    source_before = save_to_bytes(source)
    target = dest.slide_layouts[3]
    _set_content_slots(
        target,
        [(PP_PLACEHOLDER.OBJECT, 3), (PP_PLACEHOLDER.OBJECT, 2)],
    )
    before = save_to_bytes(dest)

    report = dest.import_slide(
        source,
        1,
        mode="adopt_theme",
        target_layout=target,
        placeholder_map={1: 2},
    )

    assert report.placeholder_map_used == ((0, 0), (1, 2))
    after = save_to_bytes(dest)
    assert_changed_parts(
        before,
        after,
        expect_added=["ppt/slides/_rels/slide4.xml.rels", "ppt/slides/slide4.xml"],
        expect_changed=[
            "[Content_Types].xml",
            "ppt/_rels/presentation.xml.rels",
            "ppt/presentation.xml",
        ],
    )
    _assert_clean(after)
    reopened = Presentation(io.BytesIO(after))
    imported = reopened.slides[-1]
    assert imported.slide_layout.name == "Two Content"
    placeholders = {
        shape.element.ph_idx: shape for shape in imported.shapes if shape.is_placeholder
    }
    assert set(placeholders) == {0, 2}
    assert placeholders[0].text_frame.text == "Beta Content"
    assert "Beta point one" in placeholders[2].text_frame.text
    assert zip_member_map(save_to_bytes(source)) == zip_member_map(source_before)


def test_adopt_theme_explicit_none_bakes_orphan_and_reports_null_target():
    dest = _open(ALPHA)
    source = _open(BETA)

    report = dest.import_slide(
        source, 1, mode="adopt_theme", placeholder_map={1: None}
    )

    assert report.placeholder_map_used == ((0, 0), (1, None))
    assert {"source_idx": 1, "target_idx": None} in report.to_dict()[
        "placeholder_map_used"
    ]
    reopened = save_reopen(dest)
    imported = reopened.slides[-1]
    body = next(shape for shape in imported.shapes if shape.name == "Content Placeholder 2")
    assert not body.is_placeholder
    assert "Beta point one" in body.text_frame.text


def test_adopt_theme_automatic_orphan_reports_null_target():
    dest = _open(ALPHA)

    report = dest.import_slide(
        _open(BETA),
        1,
        mode="adopt_theme",
        target_layout=dest.slide_layouts[5],  # -- Title Only has no content slot
    )

    assert report.placeholder_map_used == ((0, 0), (1, None))
    assert report.to_dict()["placeholder_map_used"] == [
        {"source_idx": 0, "target_idx": 0},
        {"source_idx": 1, "target_idx": None},
    ]


@pytest.mark.parametrize("mode", ["keep_appearance", "bake"])
def test_placeholder_map_is_rejected_for_non_reconciling_import_modes(mode):
    dest = _open(ALPHA)
    source = _open(BETA)

    error = assert_refusal_atomic(
        dest,
        lambda prs: prs.import_slide(source, 0, mode=mode, placeholder_map={0: 0}),
        ValueError,
    )

    assert "only to mode='adopt_theme'" in str(error)


@pytest.mark.parametrize(
    ("placeholder_map", "message"),
    [
        ({99: 1}, "not a placeholder on this slide"),
        ({1: 99}, "not a placeholder on the target"),
        ({0: 2, 1: 2}, "one target"),
    ],
)
def test_adopt_theme_placeholder_map_validation_is_atomic(placeholder_map, message):
    dest = _open(ALPHA)
    source = _open(BETA)

    error = assert_refusal_atomic(
        dest,
        lambda prs: prs.import_slide(
            source,
            1,
            mode="adopt_theme",
            target_layout=prs.slide_layouts[3],
            placeholder_map=placeholder_map,
        ),
        ValueError,
    )

    assert message in str(error)


def test_adopt_theme_falls_back_to_type_match_for_renamed_layout():
    """Beta's chart slide sits on 'Beta Special' - a name alpha lacks - but the layout
    kept type="titleOnly", so auto-selection falls back to alpha's Title Only."""
    dest = _open(ALPHA)
    report = dest.import_slide(_open(BETA), 2, mode="adopt_theme")
    assert report.layout_binding_method == "type-match"
    assert report.layout_binding == "/ppt/slideLayouts/slideLayout6.xml"
    reopened = save_reopen(dest)
    assert reopened.slides[-1].slide_layout.name == "Title Only"


def test_adopt_theme_unmatched_layout_refuses_and_explicit_target_recovers():
    dest = _open(ALPHA)
    source = _open(BETA)
    # -- force a truly unmatchable source layout: alien name AND no type token
    source_layout = source.slides[2].slide_layout
    source_layout._element.attrib.pop("type", None)

    before = save_to_bytes(dest)
    with pytest.raises(UnsupportedStructureError, match="target_layout"):
        dest.import_slide(source, 2, mode="adopt_theme")
    assert_changed_parts(before, save_to_bytes(dest))  # -- empty budget

    report = dest.import_slide(source, 2, mode="adopt_theme", target_layout=dest.slide_layouts[5])
    assert report.layout_binding_method == "explicit"
    saved = save_to_bytes(dest)
    _assert_clean(saved)
    reopened = Presentation(io.BytesIO(saved))
    assert reopened.slides[3].shapes.chart_by_name("beta_chart") is not None


def test_bake_freezes_look_without_importing_masters():
    dest = _open(ALPHA)
    report = dest.import_slide(_open(BETA), 1, mode="bake")
    assert report.run_shifts == ()  # -- baked: resolution cannot shift
    assert set(report.baked_shapes) == {"Title 1", "Content Placeholder 2"}

    saved = save_to_bytes(dest)
    _assert_clean(saved)
    reopened = Presentation(io.BytesIO(saved))
    assert len(reopened.slide_masters) == 1
    imported = reopened.slides[3]
    assert all(not s.is_placeholder for s in imported.shapes)  # -- all free shapes now
    title = next(s for s in imported.shapes if s.name == "Title 1")
    run = title.text_frame.paragraphs[0].runs[0]
    assert run.font.name == "Courier New"  # -- beta's look, made local
    assert run.font.size is not None


def test_bake_uses_unique_blank_only_after_name_and_type_are_absent():
    dest = _open(ALPHA)
    source = _open(BETA)
    source_layout = source.slides[2].slide_layout
    source_layout.name = None
    source_layout._element.attrib.pop("type", None)

    report = dest.import_slide(source, 2, mode="bake")

    assert report.layout_binding_method == "blank-fallback"
    reopened = save_reopen(dest)
    assert reopened.slides[-1].slide_layout.name == "Blank"
    assert reopened.slides[-1].shapes.chart_by_name("beta_chart") is not None


# ----------------------------------------------------------- unique layout selection


def test_strict_layout_name_lookup_refuses_duplicates_without_changing_first_match_lookup():
    dest = _open(ALPHA)
    first = dest.slide_layouts[0]
    duplicate = dest.slide_layouts[1]
    duplicate.name = first.name

    assert dest.slide_layouts.get_by_name(first.name) is first
    with pytest.raises(AmbiguousTargetError) as exc_info:
        dest.slide_layouts.get_unique_by_name(first.name)

    message = str(exc_info.value)
    assert "0 (%s)" % first.part.partname in message
    assert "1 (%s)" % duplicate.part.partname in message
    assert "explicitly by index or partname" in message


def test_duplicate_layout_name_on_one_master_refuses_and_explicit_target_recovers():
    dest = _open(ALPHA)
    source = _open(BETA)
    first = dest.slide_layouts[0]
    duplicate = dest.slide_layouts[1]
    duplicate.name = first.name

    error = assert_refusal_atomic(
        dest,
        lambda prs: prs.import_slide(source, 0, mode="adopt_theme"),
        AmbiguousTargetError,
    )

    _assert_layout_candidates_reported(error, (first, duplicate))
    imported = _assert_explicit_layout_survives_reopen(dest, source, 0, "adopt_theme", first)
    assert imported.shapes.title.text_frame.text == "Beta Overview"


def test_duplicate_layout_name_across_masters_refuses_atomically():
    dest = _open(ALPHA)
    dest.import_slide(_open(BETA), 2, mode="keep_appearance")
    source = _open(BETA)
    first = dest.slide_masters[0].slide_layouts[0]
    duplicate = dest.slide_masters[1].slide_layouts[0]
    duplicate.name = first.name

    error = assert_refusal_atomic(
        dest,
        lambda prs: prs.import_slide(source, 0, mode="adopt_theme"),
        AmbiguousTargetError,
    )

    _assert_layout_candidates_reported(error, (first, duplicate))


def test_duplicate_layout_type_refuses_before_unique_blank_and_explicit_target_recovers():
    dest = _open(ALPHA)
    source = _open(BETA)
    target = dest.slide_layouts[5]
    duplicate = dest.slide_layouts[4]
    duplicate._element.set("type", target._element.get("type"))

    error = assert_refusal_atomic(
        dest,
        lambda prs: prs.import_slide(source, 2, mode="bake"),
        AmbiguousTargetError,
    )

    _assert_layout_candidates_reported(error, (target, duplicate))
    imported = _assert_explicit_layout_survives_reopen(dest, source, 2, "bake", target)
    assert imported.shapes.chart_by_name("beta_chart") is not None


def test_duplicate_blank_layout_refuses_and_explicit_target_recovers():
    dest = _open(ALPHA)
    source = _open(BETA)
    source_layout = source.slides[2].slide_layout
    source_layout.name = None
    source_layout._element.attrib.pop("type", None)
    target = dest.slide_layouts[6]
    duplicate = dest.slide_layouts[5]
    duplicate._element.set("type", "blank")

    error = assert_refusal_atomic(
        dest,
        lambda prs: prs.import_slide(source, 2, mode="bake"),
        AmbiguousTargetError,
    )

    _assert_layout_candidates_reported(error, (target, duplicate))
    imported = _assert_explicit_layout_survives_reopen(dest, source, 2, "bake", target)
    assert imported.shapes.chart_by_name("beta_chart") is not None


def test_layout_ambiguity_diagnostic_is_independent_of_collection_order():
    def ambiguous_error(reorder):
        dest = _open(ALPHA)
        source = _open(BETA)
        dest.slide_layouts[1].name = dest.slide_layouts[0].name
        if reorder:
            id_list = dest.slide_masters[0].slide_layouts._sldLayoutIdLst
            first_entry = id_list.sldLayoutId_lst[0]
            id_list.remove(first_entry)
            id_list.append(first_entry)
        return str(
            assert_refusal_atomic(
                dest,
                lambda prs: prs.import_slide(source, 0, mode="adopt_theme"),
                AmbiguousTargetError,
            )
        )

    assert ambiguous_error(reorder=False) == ambiguous_error(reorder=True)


def test_keep_appearance_ignores_ambiguous_destination_layouts():
    dest = _open(ALPHA)
    source = _open(BETA)
    dest.slide_layouts[1].name = dest.slide_layouts[0].name

    report = dest.import_slide(source, 0, mode="keep_appearance")

    assert report.layout_binding_method == "transplant"
    reopened = save_reopen(dest)
    assert reopened.slides[-1].shapes.title.text_frame.text == "Beta Overview"


def test_bake_with_no_enrolled_destination_layout_refuses_typed_not_index_error():
    dest = Presentation()
    for layout in list(dest.slide_layouts):
        dest.slide_layouts.remove(layout)
    source = Presentation()
    source.slides.add_slide(source.slide_layouts[0])

    error = assert_refusal_atomic(
        dest,
        lambda prs: prs.import_slide(source, 0, mode="bake"),
        UnsupportedStructureError,
    )

    assert "target_layout" in str(error)


def test_bake_drops_furniture_placeholders():
    source = _open(GAUNTLET)
    source.apply_footers(footer="Travelling footer", slide_number=True)
    dest = _open(ALPHA)
    report = dest.import_slide(source, 0, mode="bake")
    assert len(report.dropped_placeholders) == 2  # -- ftr + sldNum
    # -- dropping shapes must not slip the shift keys: surviving runs are baked and
    # -- re-resolve identically, so there are NO shifts, phantom or real
    assert report.run_shifts == ()
    reopened = save_reopen(dest)
    imported = reopened.slides[3]
    texts = [s.text_frame.text for s in imported.shapes if s.has_text_frame]
    assert "Travelling footer" not in texts


# ------------------------------------------------------------------- cross-contamination


def test_source_is_never_mutated_and_imported_chart_is_independent():
    """Edit the imported chart; every member of the source package stays byte-identical."""
    dest = _open(ALPHA)
    source = _open(BETA)
    source_before = save_to_bytes(source)
    dest.import_slide(source, 2, mode="keep_appearance")

    chart = dest.slides[3].shapes.chart_by_name("beta_chart")
    chart.replace_data_safe(["North", "South"], [("FY26", (99.0, 1.0))])

    # -- member-by-member, never whole-file bytes: `save()` stamps wall-clock time into zip
    # -- entry headers, so two saves straddling a 2-second boundary differ in the header alone
    assert zip_member_map(save_to_bytes(source)) == zip_member_map(source_before)
    reopened_source = Presentation(io.BytesIO(save_to_bytes(source)))
    source_chart = reopened_source.slides[2].shapes.chart_by_name("beta_chart")
    values = [pt for series in source_chart.plots[0].series for pt in series.values]
    assert values == [12.5, 8.75]  # -- source data untouched


def test_import_remaps_copied_document_identity_without_mutating_source():
    from lxml import etree

    source = _open(BETA)
    dest = _open(ALPHA)
    identity_tag = "{http://schemas.microsoft.com/office/drawing/2014/main}creationId"
    source_id = "{22222222-2222-2222-2222-222222222222}"
    etree.SubElement(source.slides[0]._element, identity_tag, id=source_id)

    dest.import_slide(source, 0, mode="adopt_theme")
    imported_id = dest.slides[-1]._element.find(".//" + identity_tag).get("id")

    assert imported_id != source_id
    assert source.slides[0]._element.find(".//" + identity_tag).get("id") == source_id


def test_import_refuses_a_deleted_source_slide_proxy():
    source = _open(BETA)
    stale = source.slides[0]
    source.slides.delete(0)
    dest = _open(ALPHA)

    raised = assert_refusal_atomic(
        dest,
        lambda p: p.import_slide(source, stale, mode="adopt_theme"),
        TargetNotFoundError,
    )

    assert "source slide" in str(raised)


def test_media_always_copies_never_shared_across_packages():
    dest = _open(ALPHA)
    source = _open(BETA)
    report = dest.import_slide(source, 3, mode="adopt_theme")  # -- beta picture slide
    assert any("/ppt/media/" in part for part in report.parts_added)
    saved = save_to_bytes(dest)
    reopened = Presentation(io.BytesIO(saved))
    imported_pic = reopened.slides[3].shapes.picture_by_name("beta_pic")
    source_pic = source.slides[3].shapes.picture_by_name("beta_pic")
    assert imported_pic.image.blob == source_pic.image.blob  # -- same bytes, copied part


def test_notes_policy():
    source = _open(GAUNTLET)  # -- slide 2 has speaker notes
    dest = _open(ALPHA)
    with_notes = dest.import_slide(source, 1, mode="adopt_theme", notes=True)
    assert with_notes.notes_copied is True
    reopened = save_reopen(dest)
    assert reopened.slides[3].read_notes_text() == "Gauntlet speaker notes."

    dest2 = _open(ALPHA)
    without = dest2.import_slide(source, 1, mode="adopt_theme", notes=False)
    assert without.notes_copied is False
    reopened2 = save_reopen(dest2)
    assert not reopened2.slides[3].has_notes_slide


# ------------------------------------------------------------------ position and sections


def test_position_inserts_at_index():
    dest = _open(ALPHA)
    report = dest.import_slide(_open(BETA), 0, mode="adopt_theme", position=0)
    assert report.position == 0
    reopened = save_reopen(dest)
    assert reopened.slides[0].shapes.title.text_frame.text == "Beta Overview"
    assert len(reopened.slides) == 4


def test_unique_named_section_enrollment_persists_and_reports_identity():
    dest = _open(SECTIONS)
    source = _open(BETA)
    source_before = zip_member_map(save_to_bytes(source))
    before = save_to_bytes(dest)

    report = dest.import_slide(source, 0, mode="adopt_theme", section="Close")

    assert report.section == "Close"
    assert report.section_id == SECTION_IDS["Close"]
    after = save_to_bytes(dest)
    assert_changed_parts(
        before,
        after,
        expect_added=["ppt/slides/_rels/slide6.xml.rels", "ppt/slides/slide6.xml"],
        expect_changed=[
            "[Content_Types].xml",
            "ppt/_rels/presentation.xml.rels",
            "ppt/presentation.xml",
        ],
    )
    _assert_clean(after)
    reopened = Presentation(io.BytesIO(after))
    close_ids = _section_slide_ids(_section_by_id(reopened, SECTION_IDS["Close"]))
    assert close_ids[-1] == report.dest_slide_id
    assert reopened.slides[-1].shapes.title.text_frame.text == "Beta Overview"
    assert zip_member_map(save_to_bytes(source)) == source_before


def test_exact_section_id_selects_intended_duplicate_name_and_persists():
    dest = _open(SECTIONS)
    source = _open(BETA)
    _sections(dest)[1].set("name", "Close")

    report = dest.import_slide(
        source, 0, mode="adopt_theme", section_id=SECTION_IDS["Close"]
    )

    assert report.section == "Close"
    assert report.section_id == SECTION_IDS["Close"]
    reopened = save_reopen(dest)
    selected = _section_by_id(reopened, SECTION_IDS["Close"])
    other = _section_by_id(reopened, SECTION_IDS["Body"])
    assert _section_slide_ids(selected)[-1] == report.dest_slide_id
    assert report.dest_slide_id not in _section_slide_ids(other)
    assert reopened.slides[-1].shapes.title.text_frame.text == "Beta Overview"
    _assert_clean(save_to_bytes(dest))


def test_duplicate_section_name_refuses_atomically_with_all_candidates():
    dest = _open(SECTIONS)
    _sections(dest)[1].set("name", "Close")
    before = zip_member_map(save_to_bytes(dest))

    error = assert_refusal_atomic(
        dest,
        lambda prs: prs.import_slide(_open(BETA), 0, mode="adopt_theme", section="Close"),
        AmbiguousTargetError,
    )

    message = str(error)
    assert "section_id" in message
    assert "order=1, id=%r" % SECTION_IDS["Body"] in message
    assert "order=2, id=%r" % SECTION_IDS["Close"] in message
    assert zip_member_map(save_to_bytes(dest)) == before


def test_duplicate_section_id_refuses_atomically_instead_of_selecting_first():
    dest = _open(SECTIONS)
    _sections(dest)[1].set("id", SECTION_IDS["Intro"])
    before = zip_member_map(save_to_bytes(dest))

    error = assert_refusal_atomic(
        dest,
        lambda prs: prs.import_slide(
            _open(BETA), 0, mode="adopt_theme", section_id=SECTION_IDS["Intro"]
        ),
        AmbiguousTargetError,
    )

    message = str(error)
    assert "section=" in message
    assert "order=0, name='Intro', id=%r" % SECTION_IDS["Intro"] in message
    assert "order=1, name='Body', id=%r" % SECTION_IDS["Intro"] in message
    assert zip_member_map(save_to_bytes(dest)) == before


@pytest.mark.parametrize(
    "selector",
    [
        {"section": "No Such Section"},
        {"section_id": "{99999999-9999-4999-8999-999999999999}"},
        {"section_id": SECTION_IDS["Intro"].strip("{}")},
    ],
)
def test_missing_or_nonexact_section_selector_refuses_atomically(selector):
    dest = _open(SECTIONS)
    before = zip_member_map(save_to_bytes(dest))

    error = assert_refusal_atomic(
        dest,
        lambda prs: prs.import_slide(_open(BETA), 0, mode="adopt_theme", **selector),
        TargetNotFoundError,
    )

    assert "section" in str(error)
    assert zip_member_map(save_to_bytes(dest)) == before


@pytest.mark.parametrize(
    "selector",
    [
        {"section": 1},
        {"section_id": 1},
        {"section": "Intro", "section_id": SECTION_IDS["Intro"]},
    ],
)
def test_section_selector_argument_errors_are_atomic(selector):
    dest = _open(SECTIONS)
    before = zip_member_map(save_to_bytes(dest))

    error = assert_refusal_atomic(
        dest,
        lambda prs: prs.import_slide(_open(BETA), 0, mode="adopt_theme", **selector),
        ValueError,
    )

    assert "section" in str(error)
    assert zip_member_map(save_to_bytes(dest)) == before


def test_implicit_section_enrollment_remains_adjacent_and_reports_identity():
    dest = _open(SECTIONS)
    source = _open(BETA)

    # -- adjacent enrollment: inserting at position 1 lands in "Intro" (slide 1's
    # -- section), DIRECTLY AFTER slide 1's entry, and the report says which section
    report = dest.import_slide(source, 0, mode="adopt_theme", position=1)
    assert report.section == "Intro"  # -- actual enrollment, not the (None) argument
    assert report.section_id == SECTION_IDS["Intro"]
    reopened = save_reopen(dest)
    intro_ids = _section_slide_ids(_section_by_id(reopened, SECTION_IDS["Intro"]))
    deck_ids = [slide.slide_id for slide in reopened.slides]
    # -- section order mirrors deck order: [slide 1's id, the imported slide's id]
    assert intro_ids == [deck_ids[0], deck_ids[1]]
    _assert_clean(save_to_bytes(dest))


def test_implicit_section_enrollment_keeps_first_section_fallback():
    dest = _open(SECTIONS)
    report = dest.import_slide(_open(BETA), 0, mode="adopt_theme", position=0)

    assert report.section == "Intro"
    assert report.section_id == SECTION_IDS["Intro"]
    reopened = save_reopen(dest)
    intro_ids = _section_slide_ids(_section_by_id(reopened, SECTION_IDS["Intro"]))
    deck_ids = [slide.slide_id for slide in reopened.slides]
    assert intro_ids[:2] == deck_ids[:2]
    _assert_clean(save_to_bytes(dest))


def test_import_without_destination_sections_reports_null_section_identity():
    dest = _open(ALPHA)

    report = dest.import_slide(_open(BETA), 0, mode="adopt_theme")

    assert report.section is None
    assert report.section_id is None
    reopened = save_reopen(dest)
    assert _sections(reopened) == []


# --------------------------------------------------------------------------- append_deck


def test_append_deck_imports_all_slides_in_order():
    dest = _open(ALPHA)
    reports = dest.append_deck(_open(BETA), mode="keep_appearance")
    assert len(reports) == 4
    saved = save_to_bytes(dest)
    _assert_clean(saved)
    reopened = Presentation(io.BytesIO(saved))
    assert len(reopened.slides) == 7
    assert len(reopened.slide_masters) == 2  # -- dedupe held across the whole merge
    titles = [
        s.shapes.title.text_frame.text if s.shapes.title is not None else None
        for s in reopened.slides
    ]
    assert titles[3] == "Beta Overview"
    assert titles[5] == "Beta Chart"


def test_append_deck_validates_whole_source_before_first_write():
    """Poison the LAST source slide; the destination must stay untouched."""
    dest = _open(ALPHA)
    source = _open(BETA)
    source.slides[3].part.relate_to(
        source.slides[0].part,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide",
    )
    before = save_to_bytes(dest)
    with pytest.raises(RelationshipPolicyError):
        dest.append_deck(source, mode="keep_appearance")
    assert_changed_parts(before, save_to_bytes(dest))  # -- empty budget


def test_append_deck_later_layout_ambiguity_refuses_before_any_slide_is_added():
    dest = _open(ALPHA)
    source = _open(BETA)
    dest.slide_layouts[2].name = dest.slide_layouts[1].name
    before = save_to_bytes(dest)
    slide_count = len(dest.slides)

    error = assert_refusal_atomic(
        dest,
        lambda prs: prs.append_deck(source, mode="adopt_theme"),
        AmbiguousTargetError,
    )

    _assert_layout_candidates_reported(error, (dest.slide_layouts[1], dest.slide_layouts[2]))
    assert len(dest.slides) == slide_count
    assert_changed_parts(before, save_to_bytes(dest))  # -- empty budget


def test_append_deck_later_placeholder_ambiguity_refuses_before_any_slide_is_added():
    dest = _open(ALPHA)
    source = _open(BETA)
    target = dest.slide_layouts[1]
    _duplicate_content_slot(target, 2)
    _set_content_slots(
        target,
        [(PP_PLACEHOLDER.OBJECT, 3), (PP_PLACEHOLDER.OBJECT, 2)],
    )
    before = save_to_bytes(dest)
    slide_count = len(dest.slides)

    error = assert_refusal_atomic(
        dest,
        lambda prs: prs.append_deck(source, mode="adopt_theme"),
        AmbiguousTargetError,
    )

    assert "placeholder_map" in str(error)
    assert len(dest.slides) == slide_count
    assert_changed_parts(before, save_to_bytes(dest))


# ------------------------------------------------------------------------------- refusals


def test_refusal_ledger_unsupported_relationship():
    dest = _open(ALPHA)
    source = _open(BETA)
    source.slides[0].part.relate_to(
        source.slides[1].part,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide",
    )
    before = save_to_bytes(dest)
    with pytest.raises(RelationshipPolicyError) as excinfo:
        dest.import_slide(source, 0, mode="keep_appearance")
    assert "refusal ledger" in str(excinfo.value)
    assert isinstance(excinfo.value, PaperRefusal)
    assert_changed_parts(before, save_to_bytes(dest))  # -- empty budget


def test_argument_validation():
    dest = _open(ALPHA)
    source = _open(BETA)
    with pytest.raises(ValueError, match="mode"):
        dest.import_slide(source, 0, mode="magic")
    with pytest.raises(ValueError, match="same presentation"):
        dest.import_slide(dest, 0, mode="bake")
    with pytest.raises(ValueError, match="out of range"):
        dest.import_slide(source, 99, mode="bake")
    with pytest.raises(ValueError, match="position"):
        dest.import_slide(source, 0, mode="bake", position=99)
    with pytest.raises(ValueError, match="does not belong"):
        dest.import_slide(source, dest.slides[0], mode="bake")
    with pytest.raises(ValueError, match="target_layout does not apply"):
        dest.import_slide(source, 0, mode="keep_appearance", target_layout=dest.slide_layouts[0])
    with pytest.raises(ValueError, match="destination"):
        dest.import_slide(source, 0, mode="adopt_theme", target_layout=source.slide_layouts[0])


def test_alternate_content_refuses_for_reconciling_modes_only():
    from lxml import etree

    dest = _open(ALPHA)
    source = _open(BETA)
    spTree = source.slides[0].shapes._spTree
    etree.SubElement(
        spTree,
        "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent",
    )
    with pytest.raises(UnsupportedStructureError, match="AlternateContent"):
        dest.import_slide(source, 0, mode="adopt_theme")
    with pytest.raises(UnsupportedStructureError, match="AlternateContent"):
        dest.import_slide(source, 0, mode="bake")
    report = dest.import_slide(source, 0, mode="keep_appearance")  # -- opaque: allowed
    assert report.mode == "keep_appearance"
    _assert_clean(save_to_bytes(dest))


# ---------------------------------------------------------------- reports and determinism


def test_import_report_matches_frozen_golden():
    """Deterministic report, byte-identical to the reviewed golden."""
    dest = _open(ALPHA)
    report = dest.import_slide(_open(BETA), 0, mode="keep_appearance")
    actual = (json.dumps(report.to_dict(), indent=2, ensure_ascii=False) + "\n").encode("utf-8")
    golden_path = corpus.FIXTURES_DIR.parent / "goldens" / "import_beta_keep.import.json"
    assert actual == golden_path.read_bytes()


@pytest.mark.parametrize("mode", ["keep_appearance", "bake"])
def test_non_reconciling_reports_always_serialize_empty_placeholder_map(mode):
    report = _open(ALPHA).import_slide(_open(BETA), 0, mode=mode)

    assert report.placeholder_map_used == ()
    assert report.to_dict()["placeholder_map_used"] == []
    assert report.to_dict()["version"] == 3
    assert report.to_dict()["section"] is None
    assert report.to_dict()["section_id"] is None


def test_import_report_is_deterministic_across_runs():
    first = _open(ALPHA).import_slide(_open(BETA), 1, mode="bake").to_dict()
    second = _open(ALPHA).import_slide(_open(BETA), 1, mode="bake").to_dict()
    assert first == second


def test_import_from_libreoffice_authored_source():
    """Producer diversity: the source deck's final bytes were written by LibreOffice."""
    dest = _open(BETA)
    source = _open(LO_ALPHA)
    report = dest.import_slide(source, 1, mode="keep_appearance")
    assert report.run_shifts == ()
    saved = save_to_bytes(dest)
    _assert_clean(saved)
    reopened = Presentation(io.BytesIO(saved))
    assert len(reopened.slide_masters) == 2


def _prune_unused_furniture(prs):
    """Remove masters no slide uses, then unused layouts on the masters that remain.

    Test scaffolding, not an API: this is the caller-side pruning a delivery pipeline
    does with upstream primitives. Layouts go through upstream `SlideLayouts.remove`;
    masters need relationship surgery because upstream's `SlideMasters` is a read-only
    collection.
    """
    rId_attr = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    presentation_part = prs.part
    sldMasterIdLst = prs._element.sldMasterIdLst

    for master in list(prs.slide_masters):
        layouts = list(master.slide_layouts)
        if any(layout.used_by_slides for layout in layouts):
            for layout in layouts:
                if not layout.used_by_slides:
                    master.slide_layouts.remove(layout)
            continue
        # -- no layout of this master serves a slide: drop the master's id-list entry
        # -- and its relationship, and the whole chain becomes unreachable
        for rId, rel in list(presentation_part.rels.items()):
            if rel.is_external or rel.target_part is not master.part:
                continue
            for entry in list(sldMasterIdLst):
                if entry.get(rId_attr) == rId:
                    sldMasterIdLst.remove(entry)
            presentation_part.drop_rel(rId)
            break


def test_import_delete_prune_reimport_never_duplicates_partnames():
    """Regression: the fingerprint-dedupe cache must not
    resurrect parts that left the package - a ghost hit re-relates a part whose freed
    partname a later import reallocated, producing duplicate zip members with different
    content. The cycle below must yield a clean, fully-registered package."""
    import warnings
    import zipfile

    dest = Presentation()
    dest.slides.add_slide(dest.slide_layouts[6])
    source_alpha = _open(ALPHA)
    source_beta = _open(BETA)
    dest.import_slide(source_alpha, 0, mode="keep_appearance")
    transplanted = {
        str(p.partname)
        for p in dest.part.package.iter_parts()
        if "slideMasters/slideMaster" in str(p.partname)
    }
    dest.slides.delete(len(dest.slides) - 1)
    _prune_unused_furniture(dest)
    # -- the prune must genuinely evict the transplanted chain, or the cache below is
    # -- never asked to distinguish a live part from a ghost and the test proves nothing
    surviving = {
        str(p.partname)
        for p in dest.part.package.iter_parts()
        if "slideMasters/slideMaster" in str(p.partname)
    }
    assert surviving < transplanted
    dest.import_slide(source_beta, 0, mode="keep_appearance")
    dest.import_slide(source_alpha, 0, mode="keep_appearance")

    buf = io.BytesIO()
    with warnings.catch_warnings():
        warnings.simplefilter("error")  # -- zipfile's 'Duplicate name' warning = failure
        dest.save(buf)
    saved = buf.getvalue()
    names = zipfile.ZipFile(io.BytesIO(saved)).namelist()
    assert len(names) == len(set(names))
    partnames = [str(p.partname) for p in dest.part.package.iter_parts()]
    assert len(partnames) == len(set(partnames))
    _assert_clean(saved)
    reopened = Presentation(io.BytesIO(saved))
    reachable_masters = {
        str(p.partname)
        for p in reopened.part.package.iter_parts()
        if "slideMasters/slideMaster" in str(p.partname) and str(p.partname).endswith(".xml")
    }
    assert len(reachable_masters) == len(reopened.slide_masters)


def test_import_placeholder_picture_slide_under_reconciling_modes():
    """Regression: a placeholder PICTURE has no text frame; bake and
    adopt_theme must not crash on it."""
    source = Presentation()
    layout = source.slide_layouts[8]  # -- "Picture with Caption"
    slide = source.slides.add_slide(layout)
    slide.shapes.title.text_frame.paragraphs[0].add_run().text = "Pic ph source"
    picture_ph = next(
        ph for ph in slide.placeholders if ph.placeholder_format.type.name == "PICTURE"
    )
    from PIL import Image as PILImage

    png = io.BytesIO()
    PILImage.new("RGB", (16, 16), (5, 50, 100)).save(png, format="PNG")
    picture_ph.insert_picture(io.BytesIO(png.getvalue()))

    from pptx.shapes.picture import Picture

    for mode in ("bake", "adopt_theme"):
        dest = _open(ALPHA)
        report = dest.import_slide(source, 0, mode=mode)
        assert report.mode == mode
        saved = save_to_bytes(dest)
        _assert_clean(saved)
        reopened = Presentation(io.BytesIO(saved))
        blobs = [
            shape.image.blob
            for shape in reopened.slides[3].shapes
            if isinstance(shape, Picture) or (shape.is_placeholder and hasattr(shape, "image"))
        ]
        assert any(blob == png.getvalue() for blob in blobs)


def test_append_deck_corrupt_source_refuses_typed():
    dest = _open(ALPHA)
    corrupt = Presentation(str(corpus.fixture_path("self_generated/corrupt_dangling_sldid.pptx")))
    before = save_to_bytes(dest)
    with pytest.raises(UnsupportedStructureError, match="relationship graph is broken"):
        dest.append_deck(corrupt, mode="bake")
    assert_changed_parts(before, save_to_bytes(dest))  # -- empty budget


def test_notes_import_enrolls_destination_notes_master():
    """Regression: a destination without a notes master gets one created
    on notes import; it must be enrolled in p:notesMasterIdLst, not just related."""
    dest = _open(ALPHA)  # -- alpha has no notes, hence no notes master
    assert dest._element.notesMasterIdLst is None
    source = _open(GAUNTLET)
    dest.import_slide(source, 1, mode="adopt_theme", notes=True)
    reopened = save_reopen(dest)
    notesMasterIdLst = reopened._element.notesMasterIdLst
    assert notesMasterIdLst is not None
    entry = notesMasterIdLst.notesMasterId
    assert entry is not None
    target = reopened.part.related_part(entry.rId)
    assert "notesMaster" in str(target.partname)


# --------------------------------------------------------------------------------- lo_smoke


@pytest.mark.lo_smoke
@pytest.mark.parametrize("mode", ["adopt_theme", "keep_appearance", "bake"])
def test_imported_deck_loads_in_libreoffice(mode, tmp_path):
    dest = _open(ALPHA)
    source = _open(BETA)
    for index in range(2):
        dest.import_slide(source, index, mode=mode)
    out = tmp_path / ("import_%s.pptx" % mode)
    dest.save(str(out))
    lo_load_smoke(out, tmp_path)


@pytest.mark.parametrize("mode", ["adopt_theme", "keep_appearance", "bake"])
def test_mismatched_slide_size_refuses_atomically(mode):
    dest = _open(ALPHA)
    source = _open(BETA)
    source.slide_width = source.slide_width + 1

    raised = assert_refusal_atomic(
        dest,
        lambda p: p.import_slide(source, 0, mode=mode),
        UnsupportedStructureError,
    )
    assert "slide sizes differ" in str(raised)


def test_late_import_failure_restores_the_complete_destination(monkeypatch):
    import pptx.rebind as rebind_module

    dest = _open(ALPHA)
    source = _open(BETA)
    original = rebind_module._resolution_state
    calls = 0

    def fail_after_enrollment(slide):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise UnsupportedStructureError("forced late import failure")
        return original(slide)

    monkeypatch.setattr(rebind_module, "_resolution_state", fail_after_enrollment)
    raised = assert_refusal_atomic(
        dest,
        lambda p: p.import_slide(source, 0, mode="keep_appearance"),
        UnsupportedStructureError,
    )
    assert "forced late import failure" in str(raised)


def test_import_rewrites_relationship_ids_in_generic_xml_support_parts():
    from lxml import etree
    from PIL import Image

    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI
    from pptx.oxml.ns import qn
    from pptx.util import Inches

    source = _open("self_generated/minimal_clean.pptx")
    dest = _open("self_generated/minimal_clean.pptx")
    image_stream = io.BytesIO()
    Image.new("RGB", (2, 2), (1, 2, 3)).save(image_stream, format="PNG")
    image_stream.seek(0)
    picture = source.slides[0].shapes.add_picture(
        image_stream, Inches(1), Inches(1), Inches(1), Inches(1)
    )
    image_part = source.slides[0].part.related_part(picture._pic.blip_rId)

    generic = Part(
        PackURI("/ppt/diagrams/data99.xml"),
        "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml",
        source.part.package,
        (
            '<dgm xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/'
            'relationships" r:embed="rId2"/>'
        ).encode(),
    )
    generic.rels._add_relationship(RT.IMAGE, image_part)
    generic.rels._add_relationship(RT.IMAGE, image_part)
    source.slides[0].part.relate_to(
        generic,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramData",
    )

    dest.import_slide(source, 0, mode="keep_appearance")
    copied = next(
        part
        for part in dest.part.package.iter_parts()
        if str(part.partname).startswith("/ppt/diagrams/data")
    )
    rewritten_rId = etree.fromstring(copied.blob).get(qn("r:embed"))
    assert rewritten_rId in copied.rels
    assert len(copied.rels) == 1


def test_support_fingerprint_preserves_relationship_target_binding():
    from pptx.compose import _fingerprint
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI

    package = _open("self_generated/minimal_clean.pptx").part.package
    first = Part(PackURI("/ppt/media/a.bin"), "application/octet-stream", package, b"A")
    second = Part(PackURI("/ppt/media/b.bin"), "application/octet-stream", package, b"B")
    blob = b'<x r:a="rId1" r:b="rId2" xmlns:r="urn:r"/>'
    left = Part(PackURI("/ppt/diagrams/left.xml"), "application/xml", package, blob)
    right = Part(PackURI("/ppt/diagrams/right.xml"), "application/xml", package, blob)
    left.rels._add_relationship(RT.IMAGE, first)
    left.rels._add_relationship(RT.IMAGE, second)
    right.rels._add_relationship(RT.IMAGE, second)
    right.rels._add_relationship(RT.IMAGE, first)

    assert _fingerprint(left) != _fingerprint(right)
