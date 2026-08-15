"""Contracts for `Presentation.batch()`, the opt-in deck-wide validation batch.

`batch()` exposes nesting that `PackageTransaction` already implemented: inner
transactions skip validation, so the deck validates once at block exit instead of once
per edit. These tests pin the properties that make that trade safe — the block still
refuses an invalid deck, and refusing still restores the package exactly.
"""

from __future__ import annotations

import asyncio
import io
import threading

import pytest

from pptx import Presentation
from pptx._transaction import PackageTransaction
from pptx.errors import BoundaryViolationError, UnsupportedStructureError
from pptx.package import patch_save
from pptx.util import Inches

from .contract import save_reopen


def _deck(slide_count=4):
    prs = Presentation()
    layout = prs.slide_layouts[6]
    for idx in range(slide_count):
        slide = prs.slides.add_slide(layout)
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        shape.text_frame.text = "SLIDE-%d" % (idx + 1)
    return prs


def _texts(prs):
    return [
        next(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame) for slide in prs.slides
    ]


def _slide_partnames(prs):
    return sorted(
        str(part.partname)
        for part in prs.part.package.iter_parts()
        if "/slides/slide" in str(part.partname)
    )


def _count_reopens(monkeypatch):
    """Return a list that gains an entry per whole-deck validation."""
    calls = []
    original = PackageTransaction._stage_and_reopen_candidate

    def counting(self):
        calls.append(1)
        return original(self)

    monkeypatch.setattr(PackageTransaction, "_stage_and_reopen_candidate", counting)
    return calls


def test_batch_validates_once_instead_of_once_per_edit(monkeypatch):
    prs = _deck()
    paragraphs = [slide.shapes[0].text_frame.paragraphs[0] for slide in prs.slides]

    calls = _count_reopens(monkeypatch)
    for paragraph in paragraphs:
        paragraph.bullet.set_character("-")
    unbatched = len(calls)

    calls.clear()
    with prs.batch():
        for paragraph in paragraphs:
            paragraph.bullet.set_character("*")
    batched = len(calls)

    assert unbatched == len(paragraphs)
    assert batched == 1


def test_batch_commits_edits_and_they_survive_save_reopen():
    prs = _deck()
    with prs.batch():
        for idx, slide in enumerate(prs.slides):
            slide.shapes[0].text_frame.text = "EDITED-%d" % idx

    assert _texts(save_reopen(prs)) == ["EDITED-0", "EDITED-1", "EDITED-2", "EDITED-3"]


def test_batch_yields_the_presentation():
    prs = _deck()
    with prs.batch() as batched:
        assert batched is prs


@pytest.mark.filterwarnings("ignore:Duplicate name:UserWarning")
def test_batch_refuses_an_invalid_deck_at_block_exit():
    """A partname collision is invisible to the individual edit and fatal to the package.

    Staging the candidate writes both parts under one name, so stdlib `zipfile` warns
    before the guarded reader turns the duplicate member into the typed refusal.
    """
    prs = _deck()
    before = _texts(prs)

    with pytest.raises(UnsupportedStructureError):
        with prs.batch():
            prs.slides[0].shapes[0].text_frame.text = "GOOD-EDIT"
            slide_parts = [
                part
                for part in prs.part.package.iter_parts()
                if "/slides/slide" in str(part.partname)
            ]
            slide_parts[1].partname = slide_parts[0].partname

    assert _texts(prs) == before
    assert _slide_partnames(prs) == _slide_partnames(_deck())
    assert _texts(save_reopen(prs)) == before


def test_batch_refusal_discards_every_edit_in_the_block():
    """Documented trade: the block, not the edit, is the unit of rollback."""
    prs = _deck()
    before = _texts(prs)

    with pytest.raises(UnsupportedStructureError):
        with prs.batch():
            prs.slides[0].shapes[0].text_frame.text = "FIRST"
            prs.slides[1].shapes[0].text_frame.text = "SECOND"
            duplicated = list(prs.part.rels._rels.values())[-1]
            prs.part.rels._rels["rId1"] = duplicated

    assert _texts(prs) == before


def test_caller_exception_inside_batch_rolls_back_and_propagates():
    prs = _deck()
    before = _texts(prs)

    with pytest.raises(ValueError, match="caller failed"):
        with prs.batch():
            prs.slides[0].shapes[0].text_frame.text = "PARTIAL"
            raise ValueError("caller failed")

    assert _texts(prs) == before
    assert _texts(save_reopen(prs)) == before


@pytest.mark.filterwarnings("ignore:Duplicate name:UserWarning")
def test_batch_catches_corruption_from_operations_that_lack_their_own_transaction():
    """`Slides.add_slide` runs no transaction of its own; the block covers it anyway.

    Deleting a slide leaves a gap in the slide partname sequence, and the next
    `add_slide` reuses a partname that is still in use. Outside a batch nothing refuses
    and `save()` writes a package that cannot be reopened.
    """
    prs = _deck(slide_count=5)
    before = _texts(prs)

    with pytest.raises(UnsupportedStructureError):
        with prs.batch():
            prs.slides.delete(prs.slides[1])
            prs.slides.add_slide(prs.slide_layouts[6])

    assert _texts(prs) == before
    assert _texts(save_reopen(prs)) == before


def test_save_inside_a_block_refuses_and_writes_nothing(tmp_path):
    """Edits in an open block are unvalidated and may still roll back."""
    prs = _deck()
    destination = tmp_path / "out.pptx"

    with pytest.raises(BoundaryViolationError, match="open batch block"):
        with prs.batch():
            prs.slides[0].shapes[0].text_frame.text = "UNVALIDATED"
            prs.save(str(destination))

    assert not destination.exists()


def test_save_inside_a_nested_block_also_refuses():
    prs = _deck()
    with pytest.raises(BoundaryViolationError):
        with prs.batch():
            with prs.batch():
                prs.save(io.BytesIO())


def test_patch_save_inherits_the_block_guard(tmp_path):
    """`patch_save` reaches the package through `Presentation.save`."""
    source = tmp_path / "source.pptx"
    _deck().save(str(source))
    prs = Presentation(str(source))

    with pytest.raises(BoundaryViolationError):
        with prs.batch():
            patch_save(str(source), prs, str(tmp_path / "patched.pptx"))

    assert not (tmp_path / "patched.pptx").exists()


def test_read_only_deck_diff_still_works_inside_a_block(tmp_path):
    """`diff_decks` serializes for comparison, not to publish; the guard must not fire."""
    from pptx.diff import diff_decks

    baseline = tmp_path / "baseline.pptx"
    _deck().save(str(baseline))
    prs = Presentation(str(baseline))

    with prs.batch():
        prs.slides[0].shapes[0].text_frame.text = "CHANGED"
        deltas = diff_decks(str(baseline), prs)

    assert deltas


def test_save_after_a_block_closes_is_unaffected():
    prs = _deck()
    with prs.batch():
        prs.slides[0].shapes[0].text_frame.text = "COMMITTED"

    assert _texts(save_reopen(prs))[0] == "COMMITTED"


def test_a_block_on_one_deck_does_not_block_saving_another():
    """The guard keys on package identity, not on any transaction being open."""
    held, other = _deck(), _deck()
    with held.batch():
        other.slides[0].shapes[0].text_frame.text = "INDEPENDENT"
        buffer = io.BytesIO()
        other.save(buffer)

    assert buffer.getbuffer().nbytes > 0


def test_blocks_in_separate_threads_are_independent():
    """`_ACTIVE_TRANSACTIONS` is a ContextVar; a new thread starts with an empty stack."""
    failures: list[BaseException] = []
    entered = threading.Barrier(3, timeout=30)

    def edit_own_deck(tag: str) -> None:
        try:
            prs = _deck()
            with prs.batch():
                entered.wait()  # all three inside a block at once
                for idx, slide in enumerate(prs.slides):
                    slide.shapes[0].text_frame.text = "%s-%d" % (tag, idx)
            assert _texts(save_reopen(prs)) == ["%s-%d" % (tag, i) for i in range(4)]
        except BaseException as error:  # noqa: BLE001 -- reported, not swallowed
            failures.append(error)

    workers = [
        threading.Thread(target=edit_own_deck, args=("T%d" % i,), name="T%d" % i) for i in range(3)
    ]
    for worker in workers:
        worker.start()
    for worker in workers:
        worker.join()

    assert failures == []


def test_a_block_survives_awaiting_inside_it():
    """A block spans awaits within one task; the context stack travels with the task."""

    async def edit(tag: str) -> list[str]:
        prs = _deck()
        with prs.batch():
            for idx, slide in enumerate(prs.slides):
                slide.shapes[0].text_frame.text = "%s-%d" % (tag, idx)
                await asyncio.sleep(0)
        return _texts(save_reopen(prs))

    async def both() -> list[list[str]]:
        return list(await asyncio.gather(edit("A"), edit("B")))

    first, second = asyncio.run(both())

    assert first == ["A-%d" % i for i in range(4)]
    assert second == ["B-%d" % i for i in range(4)]


def test_batches_nest_and_only_the_outermost_validates(monkeypatch):
    prs = _deck()
    calls = _count_reopens(monkeypatch)

    with prs.batch():
        prs.slides[0].shapes[0].text_frame.text = "OUTER"
        with prs.batch():
            prs.slides[1].shapes[0].text_frame.text = "INNER"

    assert len(calls) == 1
    assert _texts(save_reopen(prs))[:2] == ["OUTER", "INNER"]


def test_batch_is_reentrant_across_sequential_blocks(monkeypatch):
    prs = _deck()
    calls = _count_reopens(monkeypatch)

    with prs.batch():
        prs.slides[0].shapes[0].text_frame.text = "ONE"
    with prs.batch():
        prs.slides[1].shapes[0].text_frame.text = "TWO"

    assert len(calls) == 2
    assert _texts(save_reopen(prs))[:2] == ["ONE", "TWO"]
