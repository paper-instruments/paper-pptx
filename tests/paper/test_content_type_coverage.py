"""Every part in a package must have a declared content type.

OPC gives each part a content type through an `Override` naming the part or a `Default`
matching its extension. A member with neither has no type at all, and PowerPoint refuses
such a package — measured in the application across every role a part can play:

    the part renamed to an undeclared extension        role                PowerPoint
    docProps/thumbnail.xyz                             never rendered      Repair
    ppt/slides/slide2.abc, its Override removed        slide 2 needs it    Repair
    ppt/media/image1.qqq                               drawn on slide 1    Repair
    ppt/media/orphan.png, nothing references it        unreachable         Repair

and the counterpart, an unreferenced part whose type IS declared, opens. So the rule keys
on physical membership, not on whether the loader would otherwise read the part. Verdicts
live in the `verifying-against-powerpoint` skill's ledger.

Neither stdlib `zipfile`, upstream python-pptx, nor LibreOffice enforces any of this;
LibreOffice renders all four shapes to identical text.
"""

from __future__ import annotations

import re
import zipfile

import pytest

from pptx import Presentation
from pptx.errors import PackageLimitError, PaperRefusal

from . import corpus

_CONTENT_TYPES = "[Content_Types].xml"


def _minimal_path():
    return corpus.fixture_path("self_generated/minimal_clean.pptx")


def _rebuild(target, *, rename=None, edit=None, extra=None):
    """Copy the minimal deck into `target`, applying one mutation."""
    source = _minimal_path()
    with zipfile.ZipFile(source) as incoming, zipfile.ZipFile(target, "w") as outgoing:
        for info in incoming.infolist():
            data = incoming.read(info.filename)
            if edit is not None:
                data = edit(info.filename, data)
            outgoing.writestr(rename(info.filename) if rename else info.filename, data)
        if extra is not None:
            extra(outgoing)
    return target


def _refuses(target, fragment):
    """Assert the typed refusal fires and leaves the file byte-identical."""
    before = target.read_bytes()
    assert issubclass(PackageLimitError, PaperRefusal)
    with pytest.raises(PackageLimitError, match=fragment):
        Presentation(target)
    assert target.read_bytes() == before


def test_an_unreferenced_member_with_no_declared_type_refuses(tmp_path):
    """The composition guard for deleting the unreachable-parts refusal.

    An unreferenced part is no longer refused for being unreferenced. This one is still
    refused, because it has no content type — the distinction PowerPoint draws, and the
    reason the two changes had to land together.
    """
    target = _rebuild(
        tmp_path / "orphan-undeclared.pptx",
        extra=lambda z: z.writestr("ppt/media/orphan.qqq", b"no content type declares this"),
    )

    _refuses(target, "orphan.qqq")
    assert "unreachable" not in _refusal_message(target)


def _refusal_message(target) -> str:
    with pytest.raises(PackageLimitError) as excinfo:
        Presentation(target)
    return str(excinfo.value)


def test_a_referenced_member_with_no_declared_type_refuses(tmp_path):
    """PowerPoint refuses this too, so reachability must not narrow the rule."""

    def rename(name):
        return "docProps/thumbnail.xyz" if name == "docProps/thumbnail.jpeg" else name

    def edit(name, data):
        return data.replace(b"thumbnail.jpeg", b"thumbnail.xyz") if name == "_rels/.rels" else data

    target = _rebuild(tmp_path / "referenced-undeclared.pptx", rename=rename, edit=edit)

    _refuses(target, "thumbnail.xyz")


def test_a_core_part_with_its_override_removed_refuses(tmp_path):
    """Slides carry an Override rather than a Default; removing it leaves no type."""

    def rename(name):
        return "ppt/slides/slide1.abc" if name == "ppt/slides/slide1.xml" else name

    def edit(name, data):
        if name == _CONTENT_TYPES:
            return re.sub(rb'<Override PartName="/ppt/slides/slide1\.xml"[^>]*/>', b"", data)
        if name == "ppt/_rels/presentation.xml.rels":
            return data.replace(b"slides/slide1.xml", b"slides/slide1.abc")
        return data

    target = _rebuild(tmp_path / "core-undeclared.pptx", rename=rename, edit=edit)

    _refuses(target, "slide1.abc")


def test_the_refusal_names_every_unresolved_member(tmp_path):
    def extra(outgoing):
        outgoing.writestr("ppt/media/one.qqq", b"first")
        outgoing.writestr("ppt/media/two.zzz", b"second")

    target = _rebuild(tmp_path / "two-undeclared.pptx", extra=extra)

    with pytest.raises(PackageLimitError) as excinfo:
        Presentation(target)
    assert "one.qqq" in str(excinfo.value)
    assert "two.zzz" in str(excinfo.value)


def test_the_refusal_states_a_remedy(tmp_path):
    target = _rebuild(
        tmp_path / "remedy.pptx",
        extra=lambda z: z.writestr("ppt/media/orphan.qqq", b"x"),
    )

    with pytest.raises(PackageLimitError) as excinfo:
        Presentation(target)
    message = str(excinfo.value)
    assert "Default" in message
    assert "Override" in message
    assert "re-save" in message


def _retyped_thumbnail(target, new_leaf, *, declaration=None):
    """Rename the referenced thumbnail, keeping the relationship pointed at it.

    Uses a referenced part rather than an added one so these cases exercise content-type
    resolution alone, independent of any rule about parts nothing points at.
    """

    def rename(name):
        return "docProps/" + new_leaf if name == "docProps/thumbnail.jpeg" else name

    def edit(name, data):
        if name == "_rels/.rels":
            return data.replace(b"thumbnail.jpeg", new_leaf.encode())
        if name == _CONTENT_TYPES and declaration is not None:
            return data.replace(b"</Types>", declaration + b"</Types>")
        return data

    return _rebuild(target, rename=rename, edit=edit)


def test_a_member_resolved_by_a_default_opens(tmp_path):
    """`.bin` is a declared Default in the minimal deck."""
    target = _retyped_thumbnail(tmp_path / "declared-default.pptx", "thumbnail.bin")

    assert len(Presentation(target).slides) == len(Presentation(_minimal_path()).slides)


def test_a_member_resolved_only_by_an_override_opens(tmp_path):
    """An extension nothing declares, rescued by an Override on the part's name."""
    target = _retyped_thumbnail(
        tmp_path / "declared-override.pptx",
        "thumbnail.qqq",
        declaration=b'<Override PartName="/docProps/thumbnail.qqq" ContentType="image/jpeg"/>',
    )

    Presentation(target)


def test_extension_matching_ignores_case(tmp_path):
    target = _retyped_thumbnail(tmp_path / "upper-extension.pptx", "thumbnail.BIN")

    Presentation(target)


def test_override_matching_ignores_case(tmp_path):
    target = _retyped_thumbnail(
        tmp_path / "override-case.pptx",
        "thumbnail.qqq",
        declaration=b'<Override PartName="/docProps/THUMBNAIL.QQQ" ContentType="image/jpeg"/>',
    )

    Presentation(target)


def test_the_content_types_part_itself_never_triggers_the_refusal(tmp_path):
    """It is not a part and carries no content type; no deck declares one for it."""
    target = _rebuild(tmp_path / "untouched.pptx")

    with zipfile.ZipFile(target) as archive:
        assert _CONTENT_TYPES in archive.namelist()
    Presentation(target)


def test_every_corpus_fixture_still_resolves():
    """A regression net: the rule must not refuse any deck the repo already ships."""
    refused = []
    for path in sorted(corpus.FIXTURES_DIR.rglob("*.pptx")):
        try:
            Presentation(path)
        except PackageLimitError as exc:  # pragma: no cover - only on regression
            if "no content type" in str(exc):
                refused.append(path.name)
        except PaperRefusal:
            continue  # -- corrupt-by-construction fixtures refuse for their own reasons
    assert refused == []
