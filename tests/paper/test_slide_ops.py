"""Contract tests: slide clone / delete / reorder / move.

This is where "opens in python-pptx but not in PowerPoint" corruption gets manufactured, so
the tests lean hardest on the oracles: exact changed-part budgets, relationship-integrity
scans on every output, LibreOffice smoke on every operation class, and cross-contamination
proofs (mutating a clone's chart leaves the original chart XML byte-identical).
"""

from __future__ import annotations

import io

import pytest
from lxml import etree

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.errors import (
    PaperRefusal,
    RelationshipPolicyError,
    TargetNotFoundError,
    UnsupportedStructureError,
)
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.package import patch_save
from pptx.slide import SlideClonePolicy
from pptx.util import Inches

from . import corpus
from .contract import (
    assert_changed_parts,
    assert_refusal_atomic,
    save_to_bytes,
    zip_member_map,
)
from .idlists import dangling_section_slide_ids, duplicate_section_slide_ids
from .lo import lo_load_smoke
from .relint import dangling_relationship_targets, missing_relationship_references

CHART_NOTES = "self_generated/chart_notes.pptx"
SHARED_MEDIA = "self_generated/shared_media.pptx"
GAUNTLET = "self_generated/gauntlet.pptx"
MINIMAL = "self_generated/minimal_clean.pptx"
LO_CHART_NOTES = "libreoffice_export/lo_chart_notes.pptx"
WALNUT_CHART_NOTES = "other_producers/walnut_chart_notes_absolute_rels.pptx"
_RELS_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def _open(relpath):
    return Presentation(str(corpus.fixture_path(relpath)))


def test_delete_refuses_when_another_slide_shares_the_notes_part():
    prs = Presentation()
    first = prs.slides.add_slide(prs.slide_layouts[6])
    second = prs.slides.add_slide(prs.slide_layouts[6])
    notes_part = first.notes_slide.part
    second.part.relate_to(notes_part, RT.NOTES_SLIDE)
    before = zip_member_map(save_to_bytes(prs))

    with pytest.raises(UnsupportedStructureError, match="notes part is shared"):
        prs.slides.delete(first)

    assert zip_member_map(save_to_bytes(prs)) == before


@pytest.mark.parametrize(("operation", "method_name"), [("move", "insert"), ("reorder", "append")])
def test_slide_reordering_rolls_back_a_late_xml_failure(monkeypatch, operation, method_name):
    prs = _open(GAUNTLET)
    slides = prs.slides
    before = zip_member_map(save_to_bytes(prs))
    list_type = type(slides._sldIdLst)
    original = getattr(list_type, method_name)
    failed = False

    def fail_after_write(element, *args):
        nonlocal failed
        result = original(element, *args)
        if element is slides._sldIdLst and not failed:
            failed = True
            raise RuntimeError("forced late slide-order failure")
        return result

    monkeypatch.setattr(list_type, method_name, fail_after_write)
    with pytest.raises(RuntimeError, match="forced late slide-order failure"):
        if operation == "move":
            slides.move(3, 0)
        else:
            slides.reorder([3, 2, 1, 0])

    assert zip_member_map(save_to_bytes(prs)) == before
    assert slides is prs.slides


def _assert_relationship_integrity(pptx_bytes):
    zip_map = zip_member_map(pptx_bytes)
    assert dangling_relationship_targets(zip_map) == []
    assert missing_relationship_references(zip_map) == []
    assert dangling_section_slide_ids(zip_map) == []
    assert duplicate_section_slide_ids(zip_map) == []


def _reopen(pptx_bytes):
    return Presentation(io.BytesIO(pptx_bytes))


# ------------------------------------------------------------------------------------ clone


def test_clone_deep_copies_chart_workbook_and_notes_with_exact_budget():
    prs = _open(CHART_NOTES)
    before = save_to_bytes(prs)
    clone = prs.slides.clone(0)
    assert prs.slides.index(clone) == 1
    after = save_to_bytes(prs)

    assert_changed_parts(
        before,
        after,
        expect_changed=[
            "[Content_Types].xml",  # -- Override entries for the new parts
            "ppt/_rels/presentation.xml.rels",
            "ppt/presentation.xml",
        ],
        expect_added=[
            "ppt/charts/_rels/chart2.xml.rels",
            "ppt/charts/chart2.xml",
            "ppt/embeddings/Microsoft_Excel_Sheet2.xlsx",
            "ppt/notesSlides/_rels/notesSlide2.xml.rels",
            "ppt/notesSlides/notesSlide2.xml",
            "ppt/slides/_rels/slide2.xml.rels",
            "ppt/slides/slide2.xml",
        ],
    )
    _assert_relationship_integrity(after)


def test_patch_save_clones_and_edits_walnut_graph_with_exact_budget(tmp_path):
    source = corpus.fixture_path(WALNUT_CHART_NOTES)
    before = source.read_bytes()
    before_map = zip_member_map(before)
    prs = Presentation(str(source))
    source_slide = prs.slides[1]
    clone = prs.slides.clone(1)
    title = next(
        shape
        for shape in clone.shapes
        if getattr(shape, "text", "")
        == "Clone this slide without sharing its chart or notes"
    )
    title.text += " — COPY"
    chart = next(shape.chart for shape in clone.shapes if shape.has_chart)
    chart_data = CategoryChartData()
    chart_data.categories = ["Q1", "Q2", "Q3", "Q4"]
    chart_data.add_series("Clone", (16, 22, 29, 38))
    chart.replace_data(chart_data)
    clone.replace_notes_text("Clone-only notes.")
    out = tmp_path / "walnut-clone.pptx"

    diff = patch_save(str(source), prs, str(out))
    after = out.read_bytes()

    assert_changed_parts(
        before,
        after,
        expect_changed=[
            "[Content_Types].xml",
            "ppt/_rels/presentation.xml.rels",
            "ppt/presentation.xml",
        ],
        expect_added=[
            "ppt/embeddings/clone-target1.xlsx",
            "ppt/notesSlides/_rels/notesSlide7.xml.rels",
            "ppt/notesSlides/notesSlide7.xml",
            "ppt/slides/_rels/slide7.xml.rels",
            "ppt/slides/charts/_rels/chart2.xml.rels",
            "ppt/slides/charts/chart2.xml",
            "ppt/slides/slide7.xml",
        ],
    )
    assert [delta.partname for delta in diff.deltas] == [
        "/[Content_Types].xml",
        "/ppt/_rels/presentation.xml.rels",
        "/ppt/embeddings/clone-target1.xlsx",
        "/ppt/notesSlides/_rels/notesSlide7.xml.rels",
        "/ppt/notesSlides/notesSlide7.xml",
        "/ppt/presentation.xml",
        "/ppt/slides/_rels/slide7.xml.rels",
        "/ppt/slides/charts/_rels/chart2.xml.rels",
        "/ppt/slides/charts/chart2.xml",
        "/ppt/slides/slide7.xml",
    ]
    reopened = Presentation(str(out))
    assert len(reopened.slides) == 7
    assert reopened.slides[2].read_notes_text() == "Clone-only notes."
    assert source_slide.read_notes_text() != "Clone-only notes."
    assert zip_member_map(after)["ppt/slides/charts/chart1.xml"] == before_map[
        "ppt/slides/charts/chart1.xml"
    ]
    _assert_relationship_integrity(after)


def test_mutating_the_clones_chart_leaves_the_original_chart_byte_identical():
    """THE cross-contamination test: the corruption class this rewrite makes impossible."""
    prs = _open(CHART_NOTES)
    prs.slides.clone(0)
    reopened = _reopen(save_to_bytes(prs))
    original_chart_xml = zip_member_map(save_to_bytes(reopened))["ppt/charts/chart1.xml"]

    clone_chart = next(s for s in reopened.slides[1].shapes if s.has_chart).chart
    chart_data = CategoryChartData()
    chart_data.categories = ["X", "Y", "Z"]
    chart_data.add_series("Mutated", (1.0, 2.0, 3.0))
    clone_chart.replace_data(chart_data)

    after_map = zip_member_map(save_to_bytes(reopened))
    assert after_map["ppt/charts/chart1.xml"] == original_chart_xml
    original_chart = next(s for s in reopened.slides[0].shapes if s.has_chart).chart
    assert [series.name for series in original_chart.series] == ["Q1", "Q2"]


def test_clone_notes_are_neither_dropped_nor_cross_linked():
    prs = _open(CHART_NOTES)
    prs.slides.clone(0)
    saved = save_to_bytes(prs)
    zip_map = zip_member_map(saved)
    for rels_member, expected_slide in (
        ("ppt/notesSlides/_rels/notesSlide1.xml.rels", "slide1.xml"),
        ("ppt/notesSlides/_rels/notesSlide2.xml.rels", "slide2.xml"),
    ):
        rels = etree.fromstring(zip_map[rels_member])
        slide_targets = [
            rel.get("Target")
            for rel in rels.iter(_RELS_NS + "Relationship")
            if rel.get("Type").endswith("/slide")
        ]
        assert len(slide_targets) == 1
        assert slide_targets[0].endswith(expected_slide), rels_member

    reopened = _reopen(saved)
    reopened.slides[1].replace_notes_text("clone-only notes")
    assert reopened.slides[0].read_notes_text() == "Speaker notes for the clone fixture."
    assert reopened.slides[1].read_notes_text() == "clone-only notes"


def test_clone_shares_media_by_default_and_copies_on_request():
    prs = _open(SHARED_MEDIA)
    prs.slides.clone(0)
    shared_map = zip_member_map(save_to_bytes(prs))
    assert [n for n in shared_map if n.startswith("ppt/media/")] == ["ppt/media/image1.png"]

    prs = _open(SHARED_MEDIA)
    prs.slides.clone(0, policy=SlideClonePolicy(share_media=False))
    copied_map = zip_member_map(save_to_bytes(prs))
    assert sorted(n for n in copied_map if n.startswith("ppt/media/")) == [
        "ppt/media/image1.png",
        "ppt/media/image2.png",
    ]
    _assert_relationship_integrity(save_to_bytes(prs))


def test_clone_can_drop_notes_by_policy_without_touching_the_source():
    prs = _open(CHART_NOTES)
    clone = prs.slides.clone(0, policy=SlideClonePolicy(deep_copy_notes=False))
    assert clone.has_notes_slide is False
    assert prs.slides[0].has_notes_slide is True
    _assert_relationship_integrity(save_to_bytes(prs))


def test_clone_copies_external_hyperlink_relationships():
    prs = _open(GAUNTLET)
    prs.slides.clone(3)  # -- the hyperlink slide
    zip_map = zip_member_map(save_to_bytes(prs))
    rels = etree.fromstring(zip_map["ppt/slides/_rels/slide5.xml.rels"])
    external = [
        rel.get("Target")
        for rel in rels.iter(_RELS_NS + "Relationship")
        if rel.get("TargetMode") == "External"
    ]
    assert external == ["https://example.com/paper"]
    _assert_relationship_integrity(save_to_bytes(prs))


def test_clone_of_a_libreoffice_chart_without_workbook_copies_style_parts():
    """LO charts carry colors/style parts and no embedded workbook; clone must cope."""
    prs = _open(LO_CHART_NOTES)
    before = save_to_bytes(prs)
    prs.slides.clone(0)
    saved = save_to_bytes(prs)
    assert_changed_parts(
        before,
        saved,
        expect_changed=[
            "[Content_Types].xml",
            "ppt/_rels/presentation.xml.rels",
            "ppt/presentation.xml",
        ],
        expect_added=[
            "ppt/charts/_rels/chart2.xml.rels",
            "ppt/charts/chart2.xml",
            "ppt/charts/colors2.xml",
            "ppt/charts/style2.xml",
            "ppt/notesSlides/_rels/notesSlide2.xml.rels",
            "ppt/notesSlides/notesSlide2.xml",
            "ppt/slides/_rels/slide2.xml.rels",
            "ppt/slides/slide2.xml",
        ],
    )
    assert not any(n.startswith("ppt/embeddings/") for n in zip_member_map(saved))
    _assert_relationship_integrity(saved)


def test_clone_after_parameter_positions_the_copy():
    prs = _open(GAUNTLET)
    clone = prs.slides.clone(0, after=2)
    clone_id = clone.slide_id
    reopened = _reopen(save_to_bytes(prs))
    assert reopened.slides[3].slide_id == clone_id

    prs = _open(GAUNTLET)
    clone = prs.slides.clone(0, after=prs.slides[1])  # -- Slide-typed after
    assert prs.slides.index(clone) == 2


def test_clone_with_two_charts_allocates_distinct_partnames():
    """Regression: parts created mid-clone are invisible to package partname allocation, so
    two deep-copied charts used to receive the SAME partname — duplicate zip members and one
    chart's data silently clobbered."""
    prs = _open(CHART_NOTES)
    chart_data = CategoryChartData()
    chart_data.categories = ["a", "b"]
    chart_data.add_series("Second", (7.0, 8.0))
    frame = prs.slides[0].shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(3), Inches(2), chart_data
    )
    frame.name = "second_chart"

    prs.slides.clone(0)
    saved = save_to_bytes(prs)
    zip_map = zip_member_map(saved)  # -- asserts no duplicate member names by itself
    charts = sorted(n for n in zip_map if n.startswith("ppt/charts/chart") and n.endswith(".xml"))
    workbooks = sorted(n for n in zip_map if n.startswith("ppt/embeddings/"))
    assert len(charts) == 4
    assert len(workbooks) == 4
    _assert_relationship_integrity(saved)

    reopened = _reopen(saved)
    clone_series = sorted(
        series.name
        for shape in reopened.slides[1].shapes
        if shape.has_chart
        for series in shape.chart.series
    )
    assert clone_series == ["Q1", "Q2", "Second"]  # -- neither chart clobbered the other


def test_clone_with_two_unshared_images_allocates_distinct_partnames():
    import io as io_module

    from PIL import Image as PILImage

    def png(color):
        buf = io_module.BytesIO()
        PILImage.new("RGB", (16, 16), color).save(buf, format="PNG")
        return buf.getvalue()

    prs = _open(MINIMAL)
    slide = prs.slides[0]
    slide.shapes.add_picture(io_module.BytesIO(png((250, 0, 0))), 0, 0, 914400)
    slide.shapes.add_picture(io_module.BytesIO(png((0, 250, 0))), 914400, 0, 914400)

    prs.slides.clone(0, policy=SlideClonePolicy(share_media=False))
    saved = save_to_bytes(prs)
    zip_map = zip_member_map(saved)
    media = sorted(n for n in zip_map if n.startswith("ppt/media/"))
    assert len(media) == 4
    assert len({zip_map[n] for n in media}) == 2  # -- two distinct images, each twice
    _assert_relationship_integrity(saved)

    reopened = _reopen(saved)
    clone_blobs = {
        s.image.blob for s in reopened.slides[1].shapes if s.shape_type.name == "PICTURE"
    }
    assert clone_blobs == {png((250, 0, 0)), png((0, 250, 0))}


def test_cloning_a_clone_and_repeated_clones_stay_consistent():
    prs = _open(CHART_NOTES)
    first_clone = prs.slides.clone(0)
    prs.slides.clone(prs.slides.index(first_clone))  # -- clone the clone
    prs.slides.clone(0)
    saved = save_to_bytes(prs)
    zip_map = zip_member_map(saved)
    assert len([n for n in zip_map if n.startswith("ppt/charts/chart")]) >= 4
    _assert_relationship_integrity(saved)
    assert len(_reopen(saved).slides) == 4


def test_clone_refuses_chart_with_unsupported_child_relationship():
    """A chart part related to something outside the allowed child set refuses atomically."""
    prs = _open(CHART_NOTES)
    chart_part = next(s for s in prs.slides[0].shapes if s.has_chart).chart.part
    image_part = prs.slides[0].part.package.get_or_add_image_part(
        io.BytesIO(
            Presentation(str(corpus.fixture_path(SHARED_MEDIA))).slides[0].shapes[0].image.blob
        )
    )
    chart_part.relate_to(image_part, "http://example.com/relationships/bogus")

    raised = assert_refusal_atomic(prs, lambda p: p.slides.clone(0), RelationshipPolicyError)
    assert "chart part" in str(raised)


def test_clone_refuses_notes_with_unsupported_child_relationship():
    prs = _open(CHART_NOTES)
    notes_part = prs.slides[0].part.part_related_by(
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide"
    )
    other_slide_part = prs.slides[0].part
    notes_part.relate_to(other_slide_part, "http://example.com/relationships/bogus")

    raised = assert_refusal_atomic(prs, lambda p: p.slides.clone(0), RelationshipPolicyError)
    assert "notes slide" in str(raised)


def test_clone_refuses_to_share_charts():
    prs = _open(CHART_NOTES)
    raised = assert_refusal_atomic(
        prs,
        lambda p: p.slides.clone(0, policy=SlideClonePolicy(deep_copy_charts=False)),
        RelationshipPolicyError,
    )
    assert "cross-contamination" in str(raised)
    assert isinstance(raised, PaperRefusal)


def test_clone_refuses_unsupported_relationship_types_atomically():
    prs = _open(MINIMAL)
    prs.slides[0].part.add_embedded_ole_object_part("Excel.Sheet.12", io.BytesIO(b"fake-ole"))
    raised = assert_refusal_atomic(prs, lambda p: p.slides.clone(0), RelationshipPolicyError)
    assert "does not support" in str(raised)


def test_clone_rejects_foreign_slides_and_bad_policy():
    prs = _open(GAUNTLET)
    other = _open(MINIMAL)
    with pytest.raises(TargetNotFoundError):
        prs.slides.clone(other.slides[0])
    with pytest.raises(ValueError):
        prs.slides.clone(0, policy="deep")


# ----------------------------------------------------------------------------------- delete


def test_delete_removes_slide_and_its_unshared_parts_with_exact_budget():
    prs = _open(CHART_NOTES)
    before = save_to_bytes(prs)
    prs.slides.delete(0)
    after = save_to_bytes(prs)
    assert_changed_parts(
        before,
        after,
        expect_changed=[
            "[Content_Types].xml",
            "ppt/_rels/presentation.xml.rels",
            "ppt/presentation.xml",
        ],
        expect_removed=[
            "ppt/charts/_rels/chart1.xml.rels",
            "ppt/charts/chart1.xml",
            "ppt/embeddings/Microsoft_Excel_Sheet1.xlsx",
            "ppt/notesSlides/_rels/notesSlide1.xml.rels",
            "ppt/notesSlides/notesSlide1.xml",
            "ppt/slides/_rels/slide1.xml.rels",
            "ppt/slides/slide1.xml",
        ],
    )
    _assert_relationship_integrity(after)
    assert len(_reopen(after).slides) == 0


def test_delete_keeps_media_shared_with_surviving_slides():
    prs = _open(SHARED_MEDIA)
    prs.slides.delete(0)
    zip_map = zip_member_map(save_to_bytes(prs))
    assert [n for n in zip_map if n.startswith("ppt/media/")] == ["ppt/media/image1.png"]
    _assert_relationship_integrity(save_to_bytes(prs))


def test_global_relationship_scan_after_every_gauntlet_delete():
    """Delete each gauntlet slide in turn; no output may carry a dangling reference."""
    for index in range(4):
        prs = _open(GAUNTLET)
        prs.slides.delete(index)
        saved = save_to_bytes(prs)
        _assert_relationship_integrity(saved)
        assert len(_reopen(saved).slides) == 3


def test_deleting_the_last_slide_leaves_a_valid_empty_deck():
    prs = _open(MINIMAL)
    prs.slides.delete(0)
    assert len(_reopen(save_to_bytes(prs)).slides) == 0


# --------------------------------------------------------------------------- reorder / move


def test_reorder_permutes_slides_and_round_trips():
    prs = _open(GAUNTLET)
    titles_before = [s.shapes.title.text if s.shapes.title else None for s in prs.slides]
    before = save_to_bytes(prs)
    prs.slides.reorder([2, 0, 3, 1])
    after = save_to_bytes(prs)
    assert_changed_parts(before, after, expect_changed=["ppt/presentation.xml"])

    titles_after = [s.shapes.title.text if s.shapes.title else None for s in _reopen(after).slides]
    assert titles_after == [titles_before[i] for i in [2, 0, 3, 1]]


@pytest.mark.parametrize("bad_order", [[0, 1, 2], [0, 1, 2, 2], [0, 1, 2, 4], []])
def test_reorder_rejects_non_permutations_atomically(bad_order):
    prs = _open(GAUNTLET)
    before = save_to_bytes(prs)
    with pytest.raises(ValueError):
        prs.slides.reorder(bad_order)
    assert_changed_parts(before, save_to_bytes(prs))


def test_move_repositions_a_single_slide():
    prs = _open(GAUNTLET)
    last = prs.slides[3]
    moved_id = last.slide_id
    before = save_to_bytes(prs)
    prs.slides.move(last, 0)
    after = save_to_bytes(prs)
    assert_changed_parts(before, after, expect_changed=["ppt/presentation.xml"])
    assert _reopen(after).slides[0].slide_id == moved_id
    with pytest.raises(ValueError):
        prs.slides.move(0, 99)
    with pytest.raises(ValueError):
        prs.slides.move(0, -1)


SECTIONS = "self_generated/sections.pptx"
_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _sections_of(pptx_bytes):
    presentation = etree.fromstring(zip_member_map(pptx_bytes)["ppt/presentation.xml"])
    return [
        (s.get("name"), [i.get("id") for i in s.iter("{%s}sldId" % _P14)])
        for s in presentation.iter("{%s}section" % _P14)
    ]


def _custom_show_rids_of(pptx_bytes):
    presentation = etree.fromstring(zip_member_map(pptx_bytes)["ppt/presentation.xml"])
    return [
        (show.get("name"), [s.get("{%s}id" % _R) for s in show.iter("{%s}sld" % _P)])
        for show in presentation.iter("{%s}custShow" % _P)
    ]


def test_delete_removes_section_membership():
    """Deleting a slide must not leave its id dangling in p14:sectionLst."""
    prs = _open(SECTIONS)
    prs.slides.delete(0)  # -- slide id 256, sole member of "Intro"
    saved = save_to_bytes(prs)
    _assert_relationship_integrity(saved)
    assert _sections_of(saved) == [
        ("Intro", []),  # -- section survives, empty
        ("Body", ["257", "258", "259"]),
        ("Close", ["260"]),
    ]
    assert len(_reopen(saved).slides) == 4


def test_delete_removes_custom_show_entries():
    prs = _open(SECTIONS)
    prs.slides.delete(1)  # -- slide id 257 = rId8, first entry of custom show "Focus"
    saved = save_to_bytes(prs)
    _assert_relationship_integrity(saved)  # -- a stale rId8 would be a missing-ref finding
    assert _custom_show_rids_of(saved) == [("Focus", ["rId10"])]


def test_clone_enrolls_copy_in_source_section_after_source():
    prs = _open(SECTIONS)
    clone = prs.slides.clone(2)  # -- slide id 258, middle of "Body"
    clone_id = str(clone.slide_id)
    saved = save_to_bytes(prs)
    _assert_relationship_integrity(saved)
    assert _sections_of(saved) == [
        ("Intro", ["256"]),
        ("Body", ["257", "258", clone_id, "259"]),
        ("Close", ["260"]),
    ]


def test_clone_does_not_enroll_in_custom_shows():
    prs = _open(SECTIONS)
    prs.slides.clone(1)  # -- source is in custom show "Focus"; the copy must not be
    saved = save_to_bytes(prs)
    _assert_relationship_integrity(saved)
    assert _custom_show_rids_of(saved) == [("Focus", ["rId8", "rId10"])]


def test_reorder_and_move_keep_section_integrity():
    prs = _open(SECTIONS)
    prs.slides.reorder([4, 3, 2, 1, 0])
    prs.slides.move(0, 2)
    saved = save_to_bytes(prs)
    _assert_relationship_integrity(saved)
    # -- sections are id-keyed: membership must be untouched by ordering operations
    assert _sections_of(saved) == [
        ("Intro", ["256"]),
        ("Body", ["257", "258", "259"]),
        ("Close", ["260"]),
    ]


def test_slide_ops_on_sectioned_deck_reopen_clean():
    prs = _open(SECTIONS)
    prs.slides.delete(4)
    prs.slides.clone(0)
    saved = save_to_bytes(prs)
    _assert_relationship_integrity(saved)
    reopened = _reopen(saved)
    assert len(reopened.slides) == 5


@pytest.mark.lo_smoke
def test_slide_ops_on_sectioned_deck_load_in_libreoffice(tmp_path):
    """The section-maintenance writes get independent-loader coverage too."""
    prs = _open(SECTIONS)
    prs.slides.delete(0)
    prs.slides.clone(1)
    prs.slides.move(0, 2)
    out = tmp_path / "sectioned_ops.pptx"
    prs.save(str(out))
    _assert_relationship_integrity(out.read_bytes())
    lo_load_smoke(out, tmp_path)


def test_delete_error_paths_leave_the_deck_untouched():
    prs = _open(GAUNTLET)
    before = save_to_bytes(prs)
    with pytest.raises(IndexError):
        prs.slides.delete(99)
    other = _open(MINIMAL)
    with pytest.raises(TargetNotFoundError):
        prs.slides.delete(other.slides[0])
    assert_changed_parts(before, save_to_bytes(prs))  # -- empty budget


def test_delete_refuses_an_additional_relationship_alias_atomically():
    prs = _open(MINIMAL)
    target = prs.slides[0]
    prs.part.rels._add_relationship(
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide",
        target.part,
    )

    raised = assert_refusal_atomic(prs, lambda p: p.slides.delete(0), PaperRefusal)

    assert "additional inbound relationship aliases" in str(raised)


def test_clone_rolls_back_late_failure_and_preserves_live_shape_proxy(monkeypatch):
    from pptx import slideops

    prs = _open(MINIMAL)
    shape = prs.slides[0].shapes[0]

    def fail_after_clone(*args):
        raise RuntimeError("forced late clone failure")

    monkeypatch.setattr(slideops, "enroll_clone_in_section", fail_after_clone)
    assert_refusal_atomic(prs, lambda p: p.slides.clone(0), RuntimeError)

    shape.text = "proxy remains live"
    assert prs.slides[0].shapes[0].text == "proxy remains live"


def test_clone_allocates_fresh_document_identity_for_each_copy():
    prs = _open(MINIMAL)
    identity_tag = "{http://schemas.microsoft.com/office/drawing/2014/main}creationId"
    etree.SubElement(
        prs.slides[0]._element,
        identity_tag,
        id="{11111111-1111-1111-1111-111111111111}",
    )

    first = prs.slides.clone(0)
    second = prs.slides.clone(0)
    ids = [slide._element.find(".//" + identity_tag).get("id") for slide in prs.slides]

    assert first is not second
    assert len(ids) == len(set(ids))


def test_layout_delete_refuses_an_additional_relationship_alias_atomically():
    prs = _open(MINIMAL)
    layout = next(layout for layout in prs.slide_layouts if not layout.used_by_slides)
    master = layout.slide_master
    master.part.rels._add_relationship(
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout",
        layout.part,
    )

    raised = assert_refusal_atomic(prs, lambda p: master.slide_layouts.remove(layout), PaperRefusal)

    assert "additional inbound relationship aliases" in str(raised)


# --------------------------------------------------------------------------------- lo_smoke


@pytest.mark.lo_smoke
@pytest.mark.parametrize("operation", ["clone", "delete", "reorder"])
def test_slide_operation_outputs_load_in_libreoffice(operation, tmp_path):
    prs = _open(GAUNTLET)
    if operation == "clone":
        prs.slides.clone(1)  # -- the chart+notes slide, the hardest case
    elif operation == "delete":
        prs.slides.delete(1)
    else:
        prs.slides.reorder([3, 2, 1, 0])
    out = tmp_path / ("%s.pptx" % operation)
    prs.save(str(out))
    lo_load_smoke(out, tmp_path)
