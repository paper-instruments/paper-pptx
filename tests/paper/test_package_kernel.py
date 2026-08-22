"""Contract tests: the package kernel (`pptx.package` additive functions).

Every pinned requirement is exercised here: the meaningful-whitespace trap
(via the frozen trailing-space fixture pair), no-op byte-identity, exact single-part budgets,
zip determinism (pinned entry order and epoch timestamps), and mid-write crash atomicity.
"""

from __future__ import annotations

import zipfile

import pytest

import pptx.package as package_module
from pptx import Presentation
from pptx.errors import PaperRefusal, UnsupportedStructureError
from pptx.package import PackageDiff, diff_package, patch_save, xml_equivalent

from . import corpus
from .contract import assert_file_bytes_unchanged, snapshot_parts, zip_member_map
from .lo import lo_load_smoke

MINIMAL = "self_generated/minimal_clean.pptx"
GAUNTLET = "self_generated/gauntlet.pptx"
PAIR_A = "self_generated/whitespace_trailing_a.pptx"
PAIR_B = "self_generated/whitespace_trailing_b.pptx"
LO_MINIMAL = "libreoffice_export/lo_minimal_clean.pptx"
WALNUT_CHART_NOTES = "other_producers/walnut_chart_notes_absolute_rels.pptx"
WALNUT_SHARED_MEDIA = "other_producers/walnut_shared_media_absolute_rels.pptx"


def _fixture(relpath):
    return str(corpus.fixture_path(relpath))


def _slide1_xml(relpath):
    return zip_member_map(corpus.fixture_path(relpath).read_bytes())["ppt/slides/slide1.xml"]


def _deck_with_extra_members(relpath, target, extras):
    """Write corpus fixture `relpath` to `target` with `extras` (name -> bytes) appended.

    Kept out of the corpus deliberately: no fixture carries these shapes today and adding
    one would change what every other suite sees.
    """
    with zipfile.ZipFile(corpus.fixture_path(relpath)) as incoming:
        with zipfile.ZipFile(str(target), "w") as outgoing:
            for info in incoming.infolist():
                outgoing.writestr(info, incoming.read(info.filename))
            for name, data in extras.items():
                outgoing.writestr(name, data)
    return target


def _folder_record_deck(target):
    """MINIMAL plus three ZIP folder records, one of them nested under another."""
    return _deck_with_extra_members(
        MINIMAL, target, {"docProps/": b"", "ppt/": b"", "ppt/slides/": b""}
    )


_CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_RELS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def _rels_xml(*relationships, root_attributes="", prefix=""):
    children = "".join(relationships)
    return (
        '<Relationships xmlns="%s"%s>%s%s</Relationships>'
        % (_RELS_NS, root_attributes, prefix, children)
    ).encode()


def _rel(rId="rId1", reltype="urn:test", target="../media/image1.png", **attributes):
    attrs = {"Id": rId, "Type": reltype, "Target": target, **attributes}
    rendered = " ".join('%s="%s"' % item for item in attrs.items())
    return "<Relationship %s/>" % rendered


def _types_xml(*declarations, root_attributes=""):
    return (
        '<Types xmlns="%s"%s>%s</Types>'
        % (_CT_NS, root_attributes, "".join(declarations))
    ).encode()


def _default(extension, content_type, extra=""):
    return '<Default Extension="%s" ContentType="%s"%s/>' % (
        extension,
        content_type,
        extra,
    )


def _override(partname, content_type, extra=""):
    return '<Override PartName="%s" ContentType="%s"%s/>' % (
        partname,
        content_type,
        extra,
    )


def _content_type_pair():
    core = "application/vnd.openxmlformats-package.core-properties+xml"
    presentation = (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
    )
    rels = "application/vnd.openxmlformats-package.relationships+xml"
    original = _types_xml(
        _default("xml", core),
        _default("rels", rels),
        _override("/ppt/presentation.xml", presentation),
    )
    serialized = _types_xml(
        _default("rels", rels),
        _default("xml", "application/xml"),
        _override("/docProps/core.xml", core),
        _override("/ppt/presentation.xml", presentation),
    )
    members = {
        "[Content_Types].xml": original,
        "_rels/.rels": _rels_xml(_rel(target="ppt/presentation.xml")),
        "docProps/core.xml": b"<core/>",
        "ppt/presentation.xml": b"<presentation/>",
    }
    candidate = dict(members)
    candidate["[Content_Types].xml"] = serialized
    return original, serialized, members, candidate


# -------------------------------------------------------------------------- xml_equivalent


def test_trailing_space_in_a_t_is_never_equivalent():
    """THE §7 trap test, on the frozen fixture pair differing by one trailing space."""
    assert not xml_equivalent(_slide1_xml(PAIR_A), _slide1_xml(PAIR_B))


def test_whitespace_only_text_of_childless_element_is_preserved():
    assert not xml_equivalent("<a><t> </t></a>", "<a><t/></a>")
    assert not xml_equivalent("<a><t> </t></a>", "<a><t>  </t></a>")


def test_structural_indentation_is_equivalent():
    """Pretty-printed vs compact (the LibreOffice-vs-python-pptx serialization difference)."""
    assert xml_equivalent('<a>\n  <b x="1">text </b>\n</a>', '<a><b x="1">text </b></a>')


def test_attribute_order_and_prefix_spelling_are_equivalent():
    assert xml_equivalent(
        '<r:x xmlns:r="urn:n" a="1" b="2">t </r:x>',
        '<q:x xmlns:q="urn:n" b="2" a="1">t </q:x>',
    )


def test_element_order_is_significant():
    assert not xml_equivalent("<a><b/><c/></a>", "<a><c/><b/></a>")


def test_malformed_xml_raises_valueerror():
    with pytest.raises(ValueError):
        xml_equivalent("not xml at all", "<a/>")


def test_identical_files_compare_equivalent_part_by_part():
    members = zip_member_map(corpus.fixture_path(MINIMAL).read_bytes())
    for name, blob in members.items():
        if name.endswith(".xml") or name.endswith(".rels"):
            assert xml_equivalent(blob, blob)


# ---------------------------------------------------------- OPC relationship equivalence


@pytest.mark.parametrize(
    "member",
    ["_rels/.rels", "ppt/slides/_rels/slide1.xml.rels"],
)
def test_relationship_bindings_ignore_child_order_and_internal_target_spelling(member):
    if member == "_rels/.rels":
        absolute = "/ppt/presentation.xml"
        relative = "ppt/presentation.xml"
    else:
        absolute = "/ppt/media/image1.png"
        relative = "../media/image1.png"
    original = _rels_xml(
        _rel("rId1", "urn:image", absolute),
        _rel("rId2", "urn:external", "https://example.com/a", TargetMode="External"),
    )
    serialized = _rels_xml(
        _rel("rId2", "urn:external", "https://example.com/a", TargetMode="External"),
        _rel("rId1", "urn:image", relative, TargetMode="Internal"),
    )

    assert package_module._members_semantically_equal(member, original, serialized)


@pytest.mark.parametrize(
    "changed",
    [
        _rel("rId9", "urn:image", "../media/image1.png"),
        _rel("rId1", "urn:other", "../media/image1.png"),
        _rel("rId1", "urn:image", "../media/image2.png"),
        _rel("rId1", "urn:image", "../media/image1.png", TargetMode="External"),
        _rel("rId1", "urn:image", "../media/image1.png", custom="changed"),
    ],
)
def test_relationship_binding_mutations_remain_significant(changed):
    original = _rels_xml(_rel("rId1", "urn:image", "../media/image1.png", custom="kept"))

    assert not package_module._members_semantically_equal(
        "ppt/slides/_rels/slide1.xml.rels", original, _rels_xml(changed)
    )


def test_relationship_root_unknown_attributes_are_compared_exactly():
    original = _rels_xml(
        _rel("rId1"), _rel("rId2"), root_attributes=' producer="walnut"'
    )
    reordered = _rels_xml(
        _rel("rId2"), _rel("rId1"), root_attributes=' producer="walnut"'
    )
    changed = _rels_xml(
        _rel("rId2"), _rel("rId1"), root_attributes=' producer="other"'
    )

    assert package_module._members_semantically_equal(
        "ppt/slides/_rels/slide1.xml.rels", original, reordered
    )
    assert not package_module._members_semantically_equal(
        "ppt/slides/_rels/slide1.xml.rels", original, changed
    )


def test_external_relationship_targets_are_not_uri_normalized():
    encoded = _rels_xml(
        _rel("rId1", "urn:external", "HTTPS://EXAMPLE.COM/a/../b", TargetMode="External")
    )
    normalized = _rels_xml(
        _rel("rId1", "urn:external", "https://example.com/b", TargetMode="External")
    )

    assert not package_module._members_semantically_equal(
        "ppt/slides/_rels/slide1.xml.rels", encoded, normalized
    )


def test_relationship_comments_are_ignored_but_processing_instructions_fall_back():
    plain = _rels_xml(_rel())
    commented = _rels_xml(_rel(), prefix="<!-- producer note -->")
    instructed = _rels_xml(_rel(), prefix="<?producer keep?>")

    assert package_module._members_semantically_equal(
        "ppt/slides/_rels/slide1.xml.rels", plain, commented
    )
    assert not package_module._members_semantically_equal(
        "ppt/slides/_rels/slide1.xml.rels", plain, instructed
    )


def test_non_xml_whitespace_relationship_text_forces_order_sensitive_fallback():
    original = _rels_xml(_rel("rId1"), _rel("rId2"), prefix="\N{NO-BREAK SPACE}")
    reordered = _rels_xml(_rel("rId2"), _rel("rId1"), prefix="\N{NO-BREAK SPACE}")

    assert not package_module._members_semantically_equal(
        "ppt/slides/_rels/slide1.xml.rels", original, reordered
    )


@pytest.mark.parametrize(
    "unsupported",
    [
        b'<Relationships xmlns="urn:wrong"><Relationship Id="rId1"/></Relationships>',
        _rels_xml('<Relationship Type="urn:test" Target="x"/>'),
        _rels_xml(_rel(TargetMode="Sideways")),
        _rels_xml(_rel(), _rel()),
        _rels_xml('<Relationship Id="rId1" Type="urn:test" Target="x">text</Relationship>'),
        _rels_xml('<Unknown/><Relationship Id="rId1" Type="urn:test" Target="x"/>'),
    ],
)
def test_unsupported_relationship_shapes_get_no_specialized_normalization(unsupported):
    assert (
        package_module._relationships_semantically_equal(
            "ppt/slides/_rels/slide1.xml.rels", unsupported, unsupported
        )
        is None
    )


def test_invalid_relationship_member_name_gets_no_specialized_normalization():
    original = _rels_xml(_rel("rId1"), _rel("rId2"))
    reordered = _rels_xml(_rel("rId2"), _rel("rId1"))

    assert not package_module._members_semantically_equal(
        "ppt/slides/slide1.xml.rels", original, reordered
    )


@pytest.mark.parametrize(
    "malformed",
    [
        b"<Relationships>",
        b'<!DOCTYPE Relationships><Relationships xmlns="%s"/>' % _RELS_NS.encode(),
        (
            b'<!DOCTYPE Relationships [<!ENTITY x "value">]>'
            b'<Relationships xmlns="%s">&x;</Relationships>' % _RELS_NS.encode()
        ),
    ],
)
def test_relationship_malformed_or_dtd_input_raises_valueerror(malformed):
    with pytest.raises(ValueError):
        package_module._members_semantically_equal(
            "ppt/slides/_rels/slide1.xml.rels", malformed, _rels_xml(_rel())
        )


# ------------------------------------------------------------- content-type equivalence


def test_content_types_compare_effective_assignments_in_package_context():
    original, serialized, members, candidate = _content_type_pair()

    assert package_module._members_semantically_equal(
        "[Content_Types].xml", original, serialized, members, candidate
    )


def test_content_type_effective_change_remains_significant():
    original, serialized, members, candidate = _content_type_pair()
    changed = serialized.replace(
        b"application/vnd.openxmlformats-package.core-properties+xml",
        b"application/x-wrong-core-properties",
    )
    candidate["[Content_Types].xml"] = changed

    assert not package_module._members_semantically_equal(
        "[Content_Types].xml", original, changed, members, candidate
    )


def test_content_type_declaration_matching_no_member_remains_significant():
    original, serialized, members, candidate = _content_type_pair()
    changed = serialized.replace(
        b"</Types>", b'<Default Extension="unused" ContentType="application/x-unused"/></Types>'
    )
    candidate["[Content_Types].xml"] = changed

    assert not package_module._members_semantically_equal(
        "[Content_Types].xml", original, changed, members, candidate
    )


@pytest.mark.parametrize(
    "change",
    [
        b'<Override PartName="/absent.xml" ContentType="application/xml"/>',
        b'<Override PartName="/docProps/core.xml" ContentType="application/xml"/>',
        b'<Default Extension="rels" ContentType="application/x-conflict"/>',
        b'<Unknown Value="x"/>',
    ],
)
def test_unsupported_content_type_shapes_fall_back_without_effective_normalization(change):
    original, serialized, members, candidate = _content_type_pair()
    changed = serialized.replace(b"</Types>", change + b"</Types>")
    candidate["[Content_Types].xml"] = changed

    assert not package_module._members_semantically_equal(
        "[Content_Types].xml", original, changed, members, candidate
    )


def test_unknown_content_type_attributes_remain_significant():
    original, serialized, members, candidate = _content_type_pair()
    changed = serialized.replace(b'<Types xmlns="', b'<Types producer="changed" xmlns="')
    candidate["[Content_Types].xml"] = changed

    assert not package_module._members_semantically_equal(
        "[Content_Types].xml", original, changed, members, candidate
    )

    declaration_changed = serialized.replace(
        b'<Default Extension="rels"', b'<Default producer="changed" Extension="rels"'
    )
    candidate["[Content_Types].xml"] = declaration_changed
    assert not package_module._members_semantically_equal(
        "[Content_Types].xml", original, declaration_changed, members, candidate
    )


def test_content_type_missing_coverage_gets_no_specialized_normalization():
    original, serialized, members, candidate = _content_type_pair()
    missing = serialized.replace(
        b'<Default Extension="xml" ContentType="application/xml"/>', b""
    ).replace(
        b'<Override PartName="/docProps/core.xml" '
        b'ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>',
        b"",
    )
    candidate["[Content_Types].xml"] = missing

    member_names = package_module._content_type_member_names(candidate)
    assert package_module._content_type_model(missing, member_names) is None
    assert not package_module._members_semantically_equal(
        "[Content_Types].xml", original, missing, members, candidate
    )


@pytest.mark.parametrize(
    "malformed",
    [
        b"<Types>",
        b'<!DOCTYPE Types><Types xmlns="%s"/>' % _CT_NS.encode(),
    ],
)
def test_content_type_malformed_or_dtd_input_raises_valueerror(malformed):
    original, _, members, candidate = _content_type_pair()
    candidate["[Content_Types].xml"] = malformed

    with pytest.raises(ValueError):
        package_module._members_semantically_equal(
            "[Content_Types].xml", original, malformed, members, candidate
        )


# --------------------------------------------------------------------------- diff_package


def test_diff_reports_exactly_the_trailing_space_part():
    diff = diff_package(_fixture(PAIR_A), _fixture(PAIR_B))
    assert [d.partname for d in diff.deltas] == ["/ppt/slides/slide1.xml"]
    assert diff.deltas[0].kind == "xml"
    assert diff.deltas[0].change == "changed"


def test_diff_of_a_package_with_itself_is_empty():
    diff = diff_package(_fixture(GAUNTLET), _fixture(GAUNTLET))
    assert diff.is_empty


@pytest.mark.parametrize("relpath", [WALNUT_CHART_NOTES, WALNUT_SHARED_MEDIA])
def test_diff_ignores_walnut_serialization_only_package_rewrites(relpath, tmp_path):
    serialized = tmp_path / "ordinary-save.pptx"
    Presentation(_fixture(relpath)).save(str(serialized))

    assert diff_package(_fixture(relpath), str(serialized)).is_empty


def test_diff_reports_added_removed_and_binary_changes(tmp_path):
    source = zip_member_map(corpus.fixture_path(MINIMAL).read_bytes())
    modified = dict(source)
    del modified["docProps/thumbnail.jpeg"]
    modified["ppt/media/new.png"] = b"png-bytes"
    modified["docProps/app.xml"] = source["docProps/app.xml"].replace(
        b"</Properties>", b"\n  </Properties>"
    )

    original_path, modified_path = tmp_path / "a.pptx", tmp_path / "b.pptx"
    for path, members in ((original_path, source), (modified_path, modified)):
        with zipfile.ZipFile(str(path), "w") as zipf:
            for name, data in members.items():
                zipf.writestr(name, data)

    diff = diff_package(str(original_path), str(modified_path))
    by_part = {d.partname: d for d in diff.deltas}
    assert by_part["/docProps/thumbnail.jpeg"].change == "removed"
    assert by_part["/ppt/media/new.png"].change == "added"
    assert by_part["/ppt/media/new.png"].kind == "binary"
    # -- the app.xml tweak only adds structural whitespace: semantically equivalent
    assert "/docProps/app.xml" not in by_part
    assert len(diff.deltas) == 2


def test_diff_to_dict_carries_pinned_schema_and_is_deterministic():
    diff = diff_package(_fixture(PAIR_A), _fixture(PAIR_B))
    payload = diff.to_dict()
    assert payload["schema"] == "paper-package-diff"
    assert payload["version"] == 1
    assert payload == diff_package(_fixture(PAIR_A), _fixture(PAIR_B)).to_dict()


# ------------------------------------------------------------------------------ patch_save


@pytest.mark.parametrize(
    "relpath", [MINIMAL, GAUNTLET, WALNUT_CHART_NOTES, WALNUT_SHARED_MEDIA]
)
def test_noop_round_trip_is_byte_identical(relpath, tmp_path):
    out = tmp_path / "noop.pptx"
    diff = patch_save(_fixture(relpath), Presentation(_fixture(relpath)), str(out))
    assert diff.is_empty
    assert out.read_bytes() == corpus.fixture_path(relpath).read_bytes()


def test_noop_round_trip_is_byte_identical_on_a_deck_carrying_folder_records(tmp_path):
    """A ZIP folder record is not a part, so `save()` structurally cannot emit one.

    Its absence from the candidate save is therefore never evidence that the document
    changed, and the documented no-op guarantee has to hold for these decks too --
    "unzip, edit, rezip" pipelines emit folder records by default.
    """
    source = _folder_record_deck(tmp_path / "folders.pptx")
    out = tmp_path / "folders_noop.pptx"

    diff = patch_save(str(source), Presentation(str(source)), str(out))

    assert diff.is_empty
    assert out.read_bytes() == source.read_bytes()
    assert len(Presentation(str(out)).slides) == 1  # -- equal bytes AND a readable deck
    assert "ppt/" in zip_member_map(out.read_bytes())


def test_edited_patch_save_changes_one_part_and_still_reports_dropped_folder_records(
    tmp_path,
):
    """An actual edit rewrites, which drops the folder records -- as `save()` and
    PowerPoint's own Save As both do. The residual must keep saying so: exactly one part
    CHANGED, and the three records honestly reported as removed rather than hidden."""
    source = _folder_record_deck(tmp_path / "folders.pptx")
    presentation = Presentation(str(source))
    presentation.slides[0].shapes.title.text_frame.paragraphs[0].runs[0].text = "Edited title"
    out = tmp_path / "folders_edit.pptx"

    diff = patch_save(str(source), presentation, str(out))

    assert [d.partname for d in diff.deltas if d.change == "changed"] == ["/ppt/slides/slide1.xml"]
    assert [d.partname for d in diff.deltas if d.change == "removed"] == [
        "/docProps/",
        "/ppt/",
        "/ppt/slides/",
    ]
    assert not [d for d in diff.deltas if d.change == "added"]
    members = zip_member_map(out.read_bytes())
    assert not [name for name in members if name.endswith("/")]
    assert Presentation(str(out)).slides[0].shapes.title.text == "Edited title"


def test_noop_patch_save_still_drops_and_reports_an_orphan_part(tmp_path):
    """The boundary the folder-record forgiveness must never widen to cover.

    An unreferenced part with a declared content type is real content: `save()` drops it
    (as PowerPoint does), so a no-op `patch_save` must rewrite rather than byte-copy, and
    must report the removal. `.jpeg` is a declared Default in MINIMAL, so the part has a
    content type and the deck is not refused at open for lacking one.
    """
    source = _deck_with_extra_members(
        MINIMAL, tmp_path / "orphan.pptx", {"ppt/media/orphan.jpeg": b"\xff\xd8\xffnot a part"}
    )
    out = tmp_path / "orphan_noop.pptx"

    diff = patch_save(str(source), Presentation(str(source)), str(out))

    assert [(d.partname, d.change) for d in diff.deltas] == [("/ppt/media/orphan.jpeg", "removed")]
    assert out.read_bytes() != source.read_bytes()
    assert "ppt/media/orphan.jpeg" not in zip_member_map(out.read_bytes())
    assert len(Presentation(str(out)).slides) == 1


def test_noop_on_a_libreoffice_file_restores_every_part_but_content_types(tmp_path):
    """Real-third-party round trip: all parts restore; the one residual delta is genuine
    (LibreOffice declares Default content types for extensions with no parts — fntdata,
    jpeg, png — and regeneration from live parts drops them)."""
    out = tmp_path / "lo_noop.pptx"
    diff = patch_save(_fixture(LO_MINIMAL), Presentation(_fixture(LO_MINIMAL)), str(out))
    assert [d.partname for d in diff.deltas] == ["/[Content_Types].xml"]


def test_single_part_edit_changes_exactly_that_part(tmp_path):
    prs = Presentation(_fixture(MINIMAL))
    prs.slides[0].shapes.title.text_frame.paragraphs[0].runs[0].text = "Edited title"
    out = tmp_path / "edit.pptx"
    diff = patch_save(_fixture(MINIMAL), prs, str(out))
    assert [d.partname for d in diff.deltas] == ["/ppt/slides/slide1.xml"]
    assert Presentation(str(out)).slides[0].shapes.title.text == "Edited title"


def test_patch_save_never_restores_over_a_trailing_space_edit(tmp_path):
    """§3's corruption-inside-the-safety-tooling scenario: removing a meaningful trailing
    space IS the edit; restore logic must not 'fix' it back to the original bytes."""
    prs = Presentation(_fixture(PAIR_A))
    box = next(s for s in prs.slides[0].shapes if s.name == "whitespace_box")
    box.text_frame.paragraphs[0].runs[0].text = "Trailing space"  # -- space removed
    out = tmp_path / "ws_edit.pptx"
    diff = patch_save(_fixture(PAIR_A), prs, str(out))
    assert [d.partname for d in diff.deltas] == ["/ppt/slides/slide1.xml"]

    reopened_box = next(
        s for s in Presentation(str(out)).slides[0].shapes if s.name == "whitespace_box"
    )
    assert reopened_box.text_frame.paragraphs[0].runs[0].text == "Trailing space"


def test_patch_save_output_is_deterministic_with_pinned_order_and_timestamps(tmp_path):
    def build():
        prs = Presentation(_fixture(MINIMAL))
        prs.slides[0].shapes.title.text_frame.paragraphs[0].runs[0].text = "Deterministic"
        return prs

    out1, out2 = tmp_path / "det1.pptx", tmp_path / "det2.pptx"
    patch_save(_fixture(MINIMAL), build(), str(out1))
    patch_save(_fixture(MINIMAL), build(), str(out2))
    assert out1.read_bytes() == out2.read_bytes()

    with zipfile.ZipFile(str(out1)) as zipf:
        names = zipf.namelist()
        assert names[0] == "[Content_Types].xml"
        assert names[1] == "_rels/.rels"
        assert names[2:] == sorted(names[2:])
        assert all(info.date_time == (1980, 1, 1, 0, 0, 0) for info in zipf.infolist())


def test_mid_write_crash_leaves_existing_output_intact(tmp_path, monkeypatch):
    """§7 failure-injection: the original survives a crash halfway through writing."""
    out = tmp_path / "out.pptx"
    precious = b"precious existing bytes"
    out.write_bytes(precious)

    prs = Presentation(_fixture(MINIMAL))
    prs.slides[0].shapes.title.text_frame.paragraphs[0].runs[0].text = "Crash"

    calls = {"count": 0}
    real_writestr = zipfile.ZipFile.writestr

    def exploding_writestr(self, *args, **kwargs):
        calls["count"] += 1
        if calls["count"] >= 5:
            raise OSError("simulated disk failure")
        return real_writestr(self, *args, **kwargs)

    monkeypatch.setattr(zipfile.ZipFile, "writestr", exploding_writestr)
    with pytest.raises(OSError, match="simulated disk failure"):
        patch_save(_fixture(MINIMAL), prs, str(out))
    monkeypatch.undo()

    assert out.read_bytes() == precious
    assert list(tmp_path.glob("*.partial")) == []  # -- no temp litter


def test_patch_save_refuses_a_non_zip_original_before_touching_anything(tmp_path):
    bogus = tmp_path / "notazip.pptx"
    bogus.write_bytes(b"this is not a zip " * 10)
    out = tmp_path / "never.pptx"
    prs = Presentation(_fixture(MINIMAL))
    before = snapshot_parts(prs)

    with assert_file_bytes_unchanged(bogus):
        with pytest.raises(UnsupportedStructureError) as exc_info:
            patch_save(str(bogus), prs, str(out))
    assert isinstance(exc_info.value, PaperRefusal)
    assert not out.exists()
    assert snapshot_parts(prs) == before


def test_patch_save_in_place_reports_the_true_residual_diff(tmp_path):
    """Regression: with out_path == original_path (in-place narrow save), the residual diff
    used to be computed AFTER the overwrite and always came back empty."""
    working = tmp_path / "deck.pptx"
    working.write_bytes(corpus.fixture_path(MINIMAL).read_bytes())

    prs = Presentation(str(working))
    prs.slides[0].shapes.title.text_frame.paragraphs[0].runs[0].text = "In-place edit"
    diff = patch_save(str(working), prs, str(working))
    assert [d.partname for d in diff.deltas] == ["/ppt/slides/slide1.xml"]
    assert Presentation(str(working)).slides[0].shapes.title.text == "In-place edit"

    # -- and an in-place no-op stays byte-identical with an empty diff
    before_bytes = working.read_bytes()
    noop_diff = patch_save(str(working), Presentation(str(working)), str(working))
    assert noop_diff.is_empty
    assert working.read_bytes() == before_bytes


def test_patch_save_rejects_a_document_that_cannot_save(tmp_path):
    with pytest.raises(ValueError):
        patch_save(_fixture(MINIMAL), object(), str(tmp_path / "x.pptx"))


def test_packagediff_type_shape():
    diff = diff_package(_fixture(MINIMAL), _fixture(MINIMAL))
    assert isinstance(diff, PackageDiff)
    assert diff.is_empty
    assert diff.to_dict()["deltas"] == []


@pytest.mark.lo_smoke
def test_patch_saved_output_loads_in_libreoffice(tmp_path):
    prs = Presentation(_fixture(MINIMAL))
    prs.slides[0].shapes.title.text_frame.paragraphs[0].runs[0].text = "LO check"
    out = tmp_path / "patched.pptx"
    patch_save(_fixture(MINIMAL), prs, str(out))
    lo_load_smoke(out, tmp_path)
