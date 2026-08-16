#!/usr/bin/env python
"""Generate deck variants for PowerPoint verification.

Each variant applies one named mutation to a control deck, so any difference in how a
reader treats them is attributable to that mutation alone. Writes the variants plus a
README the human can follow while opening them.

    uv run --no-project --with 'python-pptx==1.0.2' \
        make_variants.py OUTDIR trailing_bytes directory_entries
    uv run --no-project --with 'python-pptx==1.0.2' make_variants.py OUTDIR --all
    make_variants.py --list

Build with UPSTREAM python-pptx, never with paper-pptx: the control must be a deck the
package under test had no hand in producing.

Adding a mutation: write a function taking (raw: bytes, out: Path) and returning a
description, then register it in MUTATIONS. Keep each one minimal -- a variant that
changes two things at once cannot isolate either.
"""

from __future__ import annotations


import struct
import sys
import zipfile
from pathlib import Path
from typing import Callable, Dict, Tuple

END = struct.Struct("<4s4H2LH")
CEN = struct.Struct("<4s6H3L5H2L")
END_SIG = b"PK\x05\x06"
CEN_SIG = b"PK\x01\x02"


# --------------------------------------------------------------------------- control


def author_control(path: Path) -> None:
    """Author the unmutated deck every variant is derived from."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(8.4), Inches(1.4))
    title.text_frame.text = "PowerPoint verification"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
    title.text_frame.paragraphs[0].runs[0].font.bold = True

    body = slide.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(8.4), Inches(2.4))
    body.text_frame.word_wrap = True
    body.text_frame.text = (
        "Every file in this folder has identical slide content. They differ only in "
        "package structure. Check the window title to see which file you opened."
    )
    body.text_frame.paragraphs[0].runs[0].font.size = Pt(18)

    second = prs.slides.add_slide(prs.slide_layouts[6])
    box = second.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(8.4), Inches(2.0))
    box.text_frame.text = "Slide 2 - proves the whole package was read, not just slide 1."
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)

    prs.save(str(path))


def _end_record(raw: bytes) -> Tuple[int, list]:
    offset = raw.rfind(END_SIG)
    if offset < 0:
        raise SystemExit("control deck has no end-of-central-directory record")
    return offset, list(END.unpack_from(raw, offset))


def _central_record_starts(raw: bytes, cd_off: int, cd_size: int):
    pos, end = cd_off, cd_off + cd_size
    while pos < end:
        fields = CEN.unpack_from(raw, pos)
        if fields[0] != CEN_SIG:
            raise SystemExit("central-directory record signature mismatch")
        yield pos, fields
        pos += CEN.size + fields[10] + fields[11] + fields[12]


# --------------------------------------------------------------------------- mutations


def m_trailing_bytes(raw: bytes, out: Path) -> str:
    out.write_bytes(raw + b"--SIGNATURE-BLOCK-APPENDED-BY-SOME-TOOL--" * 3)
    return "123 bytes appended after a correct footer, as a signing or watermarking tool does"


def m_trailing_newline(raw: bytes, out: Path) -> str:
    out.write_bytes(raw + b"\n")
    return 'a single newline appended, as `echo "" >> deck.pptx` or a text-mode transfer does'


def m_bogus_comment_length(raw: bytes, out: Path) -> str:
    offset, _ = _end_record(raw)
    body = bytearray(raw)
    struct.pack_into("<H", body, offset + 20, 2560)
    out.write_bytes(bytes(body))
    return "footer declares a 2560-byte archive comment that is not present (ONE byte differs)"


def m_declared_comment(raw: bytes, out: Path) -> str:
    out.write_bytes(raw)
    with zipfile.ZipFile(out, "a") as archive:
        archive.comment = b"produced by an internal build pipeline"
    return "a real archive comment, attached the spec-correct way (footer length matches)"


def m_directory_entries(raw: bytes, out: Path) -> str:
    tmp = out.with_suffix(".src.tmp")
    tmp.write_bytes(raw)
    try:
        with zipfile.ZipFile(tmp) as incoming, zipfile.ZipFile(
            out, "w", zipfile.ZIP_DEFLATED
        ) as outgoing:
            for name in ("ppt/", "ppt/slides/", "docProps/"):
                outgoing.writestr(zipfile.ZipInfo(name), b"")
            for info in incoming.infolist():
                outgoing.writestr(info.filename, incoming.read(info.filename))
    finally:
        tmp.unlink(missing_ok=True)
    return "zero-byte 'ppt/' folder records, as Java writers and unzip-edit-rezip loops emit"


def m_prefix_data(raw: bytes, out: Path) -> str:
    offset, fields = _end_record(raw)
    prefix = b"#!/bin/sh\n# self-extracting stub\n" + b"P" * 100
    body = bytearray(prefix + raw)
    cd_size, cd_off = fields[5], fields[6]
    for pos, rec in _central_record_starts(bytes(body), cd_off + len(prefix), cd_size):
        struct.pack_into("<L", body, pos + 42, rec[16] + len(prefix))
    fields[6] = cd_off + len(prefix)
    END.pack_into(body, offset + len(prefix), *fields)
    out.write_bytes(bytes(body))
    return "a stub before the archive with every offset rebased (self-extracting shape)"


def m_stray_signature(raw: bytes, out: Path) -> str:
    stray = END_SIG + struct.pack("<4H2LH", 0, 0, 5, 5, 1234, 5678, 0)
    tmp = out.with_suffix(".src.tmp")
    tmp.write_bytes(raw)
    try:
        with zipfile.ZipFile(tmp) as incoming, zipfile.ZipFile(
            out, "w", zipfile.ZIP_DEFLATED
        ) as outgoing:
            for info in incoming.infolist():
                data = incoming.read(info.filename)
                if info.filename.endswith(".jpeg"):
                    outgoing.writestr(info.filename, data + stray, zipfile.ZIP_STORED)
                else:
                    outgoing.writestr(info.filename, data)
    finally:
        tmp.unlink(missing_ok=True)
    return "a member whose stored bytes contain the footer signature by chance"


def m_concatenated(raw: bytes, out: Path) -> str:
    out.write_bytes(raw + raw)
    return "two whole packages in one file (`cat a.pptx b.pptx`) -- genuinely ambiguous"


def m_orphan_part(raw: bytes, out: Path) -> str:
    tmp = out.with_suffix(".src.tmp")
    tmp.write_bytes(raw)
    try:
        with zipfile.ZipFile(tmp) as incoming, zipfile.ZipFile(
            out, "w", zipfile.ZIP_DEFLATED
        ) as outgoing:
            for info in incoming.infolist():
                outgoing.writestr(info.filename, incoming.read(info.filename))
            outgoing.writestr("ppt/media/orphan.png", b"\x89PNG\r\n\x1a\n" + b"\x00" * 64)
    finally:
        tmp.unlink(missing_ok=True)
    return "a part no relationship points at (upstream drops it silently on save)"


MUTATIONS: Dict[str, Callable[[bytes, Path], str]] = {
    "trailing_bytes": m_trailing_bytes,
    "trailing_newline": m_trailing_newline,
    "bogus_comment_length": m_bogus_comment_length,
    "declared_comment": m_declared_comment,
    "directory_entries": m_directory_entries,
    "prefix_data": m_prefix_data,
    "stray_signature": m_stray_signature,
    "concatenated": m_concatenated,
    "orphan_part": m_orphan_part,
}


README_HEADER = """PowerPoint verification set
===========================

Every file here was derived from 00-CONTROL-unmodified.pptx by one mutation. Slide content
is identical across all of them, so a rendering difference comes from the mutation alone.

HOW TO RUN THIS
---------------

Open each file in the real PowerPoint application and record, per file, one of:

  OPENS          - opens and looks like the control
  REPAIR         - "PowerPoint found a problem with content ... click Repair"
  OPENS-WRONG    - opens but the content differs from the control

The control must open. If it does not, something unrelated is broken and the rest of the
set tells you nothing.

Report REPAIR and OPENS-WRONG results back -- those are the ones that decide whether a
guard belongs in the package.

WHAT EACH FILE IS
-----------------
"""


def main() -> None:
    argv = sys.argv[1:]
    if not argv or argv[0] in {"-h", "--help", "--list"}:
        print("mutations:")
        for name, fn in MUTATIONS.items():
            summary = (fn.__doc__ or "").strip().splitlines()[0] if fn.__doc__ else ""
            print("  %-22s %s" % (name, summary))
        print("\nusage: make_variants.py OUTDIR [mutation ...] | --all")
        return

    outdir = Path(argv[0]).expanduser()
    requested = list(MUTATIONS) if "--all" in argv[1:] else argv[1:]
    unknown = [n for n in requested if n not in MUTATIONS]
    if unknown:
        raise SystemExit("unknown mutation(s): %s (try --list)" % ", ".join(unknown))
    if not requested:
        raise SystemExit("name at least one mutation, or pass --all")

    outdir.mkdir(parents=True, exist_ok=True)
    control = outdir / "00-CONTROL-unmodified.pptx"
    author_control(control)
    raw = control.read_bytes()

    lines = ["  00-CONTROL-unmodified.pptx", "      the control; nothing mutated", ""]
    for index, name in enumerate(requested, start=1):
        target = outdir / ("%02d-%s.pptx" % (index, name.replace("_", "-")))
        description = MUTATIONS[name](raw, target)
        differing = (
            sum(a != b for a, b in zip(raw, target.read_bytes()))
            if target.stat().st_size == len(raw)
            else None
        )
        note = "  (exactly 1 byte differs from the control)" if differing == 1 else ""
        lines += ["  %s" % target.name, "      %s%s" % (description, note), ""]
        print("  %-46s %8d bytes" % (target.name, target.stat().st_size))

    (outdir / "README.txt").write_text(README_HEADER + "\n".join(lines), encoding="utf-8")
    print("\nwrote %s" % (outdir / "README.txt"))


if __name__ == "__main__":
    main()
