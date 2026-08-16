#!/usr/bin/env python
"""Report how one pptx reader treats every deck in a folder.

Run it once per reader; the two runs together give the compatibility matrix. The two
distributions both own the `pptx` import name and refuse to coexist, so they must be run
in separate environments:

    # upstream
    uv run --no-project --with 'python-pptx==1.0.2' probe_readers.py DIR
    # the package under test
    uv run --project /path/to/paper-pptx probe_readers.py DIR

Prints one row per deck: what stdlib `zipfile` sees, then whether the reader opened it and
what text it found. Comparing the text matters as much as the open/refuse verdict -- a
reader that opens an ambiguous archive and shows the wrong content is the failure this
package exists to prevent, and it looks like success from the outside.
"""

from __future__ import annotations

import hashlib
import sys
import zipfile
from pathlib import Path


def reader_identity() -> str:
    """Identify the reader from the IMPORTED module, never from installed metadata.

    `importlib.metadata.version("paper-pptx")` can resolve from stale dist-info left
    around the tree even when the `pptx` actually imported is upstream, which silently
    mislabels the whole matrix. The fork carries `_version.__paper_version__`; upstream
    does not, so the imported module answers the question without ambiguity.
    """
    import pptx

    try:
        from pptx._version import __paper_version__

        name = "paper-pptx %s" % __paper_version__
    except ImportError:
        name = "python-pptx (upstream)"
    return "%s\n  module: %s" % (name, pptx.__file__)


def zip_view(path: Path) -> str:
    try:
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            corrupt = archive.testzip()
        return "n=%-3d crc=%s" % (len(names), "clean" if corrupt is None else "BAD")
    except Exception as exc:  # noqa: BLE001 -- reporting tool, any failure is a datum
        return "%s" % type(exc).__name__


def content_signature(path: Path) -> str:
    """Hash the visible text so two readers' interpretations can be compared."""
    from pptx import Presentation

    prs = Presentation(str(path))
    texts = [
        shape.text_frame.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text
    ]
    digest = hashlib.sha256(" ".join(texts).encode()).hexdigest()[:12]
    return "OPENS slides=%d text=%s" % (len(prs.slides), digest)


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit("usage: probe_readers.py DIR")
    folder = Path(sys.argv[1]).expanduser()
    decks = sorted(p for p in folder.iterdir() if p.suffix.lower() == ".pptx")
    if not decks:
        raise SystemExit("no .pptx files in %s" % folder)

    print("=" * 100)
    print("READER: %s" % reader_identity())
    print("=" * 100)
    for deck in decks:
        try:
            verdict = content_signature(deck)
        except Exception as exc:  # noqa: BLE001 -- a refusal is the interesting result
            verdict = "REFUSED %s: %s" % (type(exc).__name__, str(exc)[:70])
        print("%-46s %-18s %s" % (deck.name, zip_view(deck), verdict))

    print()
    print("Same text hash as the control => the reader saw the same document.")
    print("A different hash while still OPENING is the dangerous case, not the refusals.")


if __name__ == "__main__":
    main()
